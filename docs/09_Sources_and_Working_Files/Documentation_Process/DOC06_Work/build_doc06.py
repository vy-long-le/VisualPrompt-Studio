from pathlib import Path
import re, json, hashlib
from datetime import date
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[3]
WORK=Path(__file__).resolve().parent
OUT=ROOT/'06_Testing/Test_Plans/VisualPrompt_Studio_DOC-06_Test_Release_and_Maintenance_Plan_v1.0.docx'
SOURCES=[
 ('DOC-01 v2.3','00_Project_Management/Project_Overview/VisualPrompt_Studio_DOC-01_Product_and_Project_Brief_v2.3.docx'),
 ('DOC-02 v2.2','01_Requirements/Product_Requirements/VisualPrompt_Studio_DOC-02_Product_and_Functional_Requirements_Specification_v2.2.docx'),
 ('DOC-03 v2.1','02_Architecture/Data_Architecture/VisualPrompt_Studio_DOC-03_Domain_State_and_Data_Design_v2.1.docx'),
 ('DOC-04 v1.9','02_Architecture/Frontend_Architecture/VisualPrompt_Studio_DOC-04_Frontend_Architecture_and_ADRs_v1.9.docx'),
 ('DOC-05 v1.5','04_Frontend_Design/UX_and_Interaction/VisualPrompt_Studio_DOC-05_UX_and_Interaction_Specification_v1.5.docx')]

GROUPS=[]
def group(name,level,lines):
    cases=[]
    for line in lines.strip().splitlines():
        ident,refs,title,action,expected=line.split('|')
        cases.append(dict(id=ident,refs=refs,title=title,action=action,expected=expected,level=level,status='Not Run'))
    GROUPS.append((name,level,cases))

group('Catalogue và dữ liệu nghiên cứu','Unit / build-time',r'''
CAT-01|FR-001–005, FR-070, FR-071–075, FR-091|Phạm vi Workflow|Load catalogue runtime đã đối chiếu quyết định chủ dự án; thử vào mỗi card.|Chỉ Product Advertising and Photography cho phép Build trong release này; Workflow chưa nghiên cứu có trạng thái trung thực, không tạo flow giả. FR-071–075 thuộc release sau.
CAT-02|FR-010–019, FR-104|Toàn vẹn Preset/profile|Duyệt cả 10 Presets và 10 profiles; inject ID trùng, field/profile reference thiếu, default ngoài dictionary.|Bản hợp lệ qua guard; mỗi lỗi tham chiếu bị build-time validation chặn. Start Blank không được tính thành Preset thứ 11; sourcePresetId không quyết định profile.
CAT-03|FR-020–025, FR-104|30 fields và grouping|Duyệt 30 field IDs, loại control, limits, allowed values; đối chiếu mapping vào 5 nhóm Build.|Mỗi field xuất hiện đúng một lần; không mất/trùng field. Hai nhóm task-critical mở mặc định; 9 research groups không bị biến thành 9 nhóm giao diện.
CAT-04|FR-030–039, FR-104|Toàn vẹn rule|Duyệt V01–V14 và mọi fieldIds; inject severity/ID/reference không hợp lệ.|Đúng 6 errors và 8 warnings theo contract; invalid rule/reference chặn catalogue build; mọi issue có ruleId, severity, explanation, fieldIds.
CAT-05|FR-013, FR-095–098, FR-104|Encyclopedia runtime|Duyệt cả 30 field entries; search tên và thuật ngữ liên quan; kiểm tra English copy.|Mỗi field có entry khớp stable ID; tìm được nội dung tương ứng; không đưa Vietnamese research copy hoặc language switch giả vào UI.
CAT-06|FR-093–094, FR-107–109, FR-119|Preview mapping và provenance|Đối chiếu 10 Preset IDs với 6 ảnh mẫu hiện có; thay đường dẫn bằng asset thiếu.|Map bằng ID được duyệt, không suy từ tên ảnh; mỗi Preset có đúng ảnh hoặc fallback, tên/intent/action còn dùng được. Provenance nói rõ direct art-direction prompt, không chứng minh full compiler hoặc image fidelity.
CAT-07|FR-104–106, FR-112|Reconcile research → runtime|So sánh adapter với JSON archival: runtimeReady:false, approved:false, scopeHint cũ và nội dung Vietnamese.|Checklist ghi từng chuyển đổi theo baseline mới; source archival không bị sửa âm thầm. Không coi research flags là approval runtime; không thêm Admin hoặc form-builder để giải quyết reconciliation.
CAT-08|FR-030–039, FR-040–049|Port 372 fixtures|Với từng fixtureId, xác định concrete input, oracle, loại unit và expected result trước khi chạy implementation.|Fixture thiếu executable input được ghi Blocked cho test instance, không tự đánh Pass. Research actual/result không được dùng làm actual của production run; giữ mapping fixtureId → case/instance.
''')
group('PromptValidator và các cổng mutation','Unit + integration',r'''
VAL-01|FR-030–031, FR-038, FR-068|V01: dữ liệu bắt buộc|Tạo valid base; lần lượt blank/whitespace product-description, use-case và Prompt name khi tạo record.|Thiếu subject/use-case cho V01 theo contract; name blank bị record guard chặn riêng. Generate/Remix không compile hoặc ghi record khi còn error.
VAL-02|FR-032, FR-068, FR-079|V02: shape, enum, limits|Parameterize unknown field, non-string, fixed choice lạ, controls trái phép, max−1/max/max+1 Unicode scalars cho mỗi giới hạn.|Payload ngoài contract bị error; giá trị đúng boundary được chấp nhận. Không silently coerce/drop keys hoặc đếm UTF-16 code units thay Unicode scalars.
VAL-03|FR-030–039|V03: external reference|Chọn external-reference mode với reference-notes blank rồi bổ sung notes hợp lệ.|Blank tạo blocking error gắn đúng field; bổ sung notes và chạy lại loại error. Không yêu cầu upload ảnh hoặc gọi external API.
VAL-04|FR-030–039|V04: Render supplied copy|Chọn Render supplied copy; exact-copy blank/whitespace rồi nonblank.|Blank bị chặn; nonblank được xét theo các rule còn lại. Thông báo chỉ đúng copy fields, không lẫn điều kiện preserve physical label.
VAL-05|FR-030–039|V05: No added copy|Chọn No added copy với exact-copy nonblank; sau đó bỏ exact-copy.|Nonblank tạo error; blank loại error. Product label cần giữ không tự bị suy diễn thành added copy.
VAL-06|FR-030–039|V06: primary listing + supplied copy|Chọn Google Merchant primary image và Render supplied copy; đổi từng operand sang non-match.|Đúng curated match tạo error; mỗi non-match không tạo V06. Không suy diễn arbitrary free text thành match.
VAL-07|FR-033–039|V07: detail target|Dùng tight-detail/macro shot với detail-target blank; thử nonblank và non-matching shot.|Chỉ curated match + blank tạo warning; có target hoặc shot khác loại warning.
VAL-08|FR-033–039|V08: arrangement/angle|Dùng overhead flat-lay arrangement với oblique canonical angle; thay từng operand.|Match tạo warning ở các field liên quan; non-match không warning; không block bằng error.
VAL-09|FR-033–039|V09: light/shadow|Parameterize curated soft diffused/broad high-key với crisp graphic shadow và các non-matches.|Cặp mâu thuẫn tạo warning; không mở rộng semantic rule sang custom text không có trong contract.
VAL-10|FR-033–039|V10: product/copy space|Thử product left + copy left, product right + copy right; thử hai phía khác nhau.|Hai same-side matches tạo warning, opposite-side không tạo V10; issue chỉ đúng các fields.
VAL-11|FR-033–039|V11: listing/props|Primary listing với props nonblank rồi blank hoặc use-case khác.|Chỉ match tạo warning; không xóa props hoặc tự đổi use-case.
VAL-12|FR-033–039|V12: reference/preserve|External-reference mode với preserve-details blank rồi nonblank/non-match mode.|Match tạo warning; field mapping chính xác; reference-notes error vẫn ưu tiên nếu đồng thời tồn tại.
VAL-13|FR-033–039|V13: listing/copy reserve|Primary listing + Reserve space for later copy; thử policy/use-case khác.|Match tạo warning; không được coi copy-space reservation là supplied copy error.
VAL-14|FR-019, FR-033–039|V14: scale profile|Dùng product-p06-mold với scale-cue blank rồi nonblank; đổi preset fields nhưng giữ profile.|Blank ở profile p06 tạo warning; nonblank loại warning. Field customization không đổi templateProfileId nên không né V14.
VAL-15|FR-030–039, FR-084, FR-102|Errors + warnings cùng lúc|Inject ít nhất 2 errors và 2 warnings vào Generate và Remix; sửa errors rồi invoke lại.|Errors chặn compile/write dù warnings cũng có; tất cả issue được trình bày và định vị. Sau khi errors hết, warnings mới cho phép explicit confirmation.
VAL-16|FR-034–039, FR-084|Confirm, cancel và refresh|Với warnings-only, thử confirm, cancel/Escape, sửa draft trong lúc pending rồi invoke lại.|Confirm hợp lệ mới compile/save; cancel không mutation. Draft thay đổi làm confirmation cũ mất hiệu lực; validator tính lại, không dùng cached approval.
VAL-17|FR-036–039, FR-068, FR-079|Write-time revalidation|Sau bước UI validation, thay invalid config hoặc record metadata trước store write; chạy Save/Update.|Full configuration, record và next envelope được validate ngay trước write; lỗi chặn write và giữ draft/stored record cũ, không chỉ dựa validation trước đó.
''')
group('PromptCompiler và output contract','Unit',r'''
CMP-01|FR-040–042, FR-076|10 profile molds|Dùng valid base + defaults của từng p01–p10; compile với product-molds-1.0.0.|Output English raw Markdown, Balanced; sections/hierarchy/order theo profile tương ứng. Oracle là contract được duyệt, không snapshot lấy từ implementation đang kiểm thử.
CMP-02|FR-041, FR-066|Pure/deterministic|Compile cùng immutable input nhiều lần, thay clock/DOM/storage mocks bên ngoài.|Output byte-identical; không dùng DOM, clock, random, storage hoặc API. Input object không bị mutate.
CMP-03|FR-043–044|Optional omission|Lần lượt blank từng optional field và toàn bộ optional section; dùng filled counterparts.|Bỏ cả clause/section rỗng đúng contract; không để dangling heading, placeholder, null/undefined hoặc trường lặp lại.
CMP-04|FR-019, FR-040–049|Unknown contract|Inject workflow/profile/compilerVersion không hỗ trợ, field lạ, non-string và invalid fixed choice.|Reject rõ ràng trước output; không fallback sang profile gần giống. Unknown version của saved record cần reviewed migration; không thay stored generated/edited text.
CMP-05|FR-040–049|Ordinary normalization|Dùng NFC/decomposed accents, leading/trailing/multiple whitespace và mixed case ở ordinary fields.|NFC + trim + collapse whitespace theo contract; giữ case/wording, không dịch hoặc sáng tác nội dung mới.
CMP-06|FR-040–049|Exact-copy fidelity|Dùng decomposed Unicode, spaces hai đầu/nội bộ, quotes, LF/CRLF/CR, whitespace-only trong exact-copy.|Không NFC/trim/collapse exact-copy; chỉ newline CRLF/CR → LF. Whitespace-only được coi blank; JSON-quoted literal giữ nội dung và escape quotes/backslash đúng contract.
CMP-07|FR-040–049|Literal substitution|User text chứa {field}, {{nested}}, $&, $1 và backslash; compile một lần.|Placeholder substitution one-pass callback; user braces/dollar không bị interpret hoặc recursive substitution. Nội dung xuất ra đúng escaped literal.
CMP-08|FR-040–049|Markdown structural escaping|Parameterize backslash, backtick, *, _, [, ], <, >, # trong ordinary free text; thử HTML/script-like text.|Escape đúng contract, không phá cấu trúc Markdown; UI không thực thi user text. Exact-copy theo JSON literal riêng, không double-normalize.
CMP-09|FR-002–005, FR-011, FR-040–049|Start Blank neutral|Start Blank, chỉ nhập required values, chọn neutral blank profile theo contract.|Không kế thừa Preset defaults/provenance; output chỉ có required clauses và optional đã nhập. Không suy ra art-direction values ngoài input.
CMP-10|FR-040–049, FR-076|Copy policies và field-once|Parameterize ba copy-policy variants, full-filled 30 fields và optional empties.|Trusted profile order, mỗi field được emit đúng vị trí một lần; copy instruction đúng policy, không đưa field debug/CLI flags hoặc UI metadata vào prompt.
''')
group('Build, Preset và Prompt lifecycle','Integration + end-to-end',r'''
APP-01|FR-001–005, FR-010–019|Apply/override/profile|Apply từng Preset, chỉnh nhiều fixed/suggested/free fields, Generate, save/load và Remix.|Đúng structured defaults; overrides giữ nguyên đến thao tác chủ ý; templateProfileId ổn định, sourcePresetId chỉ provenance. Subject facts không được Preset tự bịa.
APP-02|FR-010–025|Start Blank/reset/stale|Apply Preset rồi Start Blank; tạo output rồi đổi config field.|Blank bỏ Preset-controlled configuration theo blank contract; sau config change output có stale indication, không tự regenerate hoặc reverse-update config.
APP-03|FR-045–049, FR-077|Manual edit/effective text|Generate rồi sửa textarea; Copy; save/load; thử không có editedPrompt.|Config bất biến bởi text edit; effectivePrompt = editedPrompt ?? generatedPrompt; Copy và persisted edited text đúng nội dung.
APP-04|FR-077|Regenerate edited|Với editedPrompt tồn tại, invoke Regenerate rồi cancel và confirm trong hai instance; thử unedited counterpart.|Cancel giữ config/output/edit/store; confirm thay output mới vào chỗ edited text, effectivePrompt bằng output mới theo AC-028/029. Unedited không hỏi mất manual edit không tồn tại.
APP-05|FR-050–054, FR-068, FR-083–084|Generate autosave|Valid draft, default No Project hoặc chọn existing/new Project; Generate.|Destination hiện rõ trước action; compile + successful write tạo Saved Prompt mới đúng membership. New Project không làm mất config. Không báo Saved trước Promise write thành công.
APP-06|FR-051–054, FR-077, FR-079|Load/Update identity|Load record, sửa config/text/name/destination hợp lệ, Update rồi reload.|Update giữ id/createdAt, tăng updatedAt sau write thành công; full record round-trip. Không tạo clone hoặc reset profile khi load.
APP-07|FR-052–054, FR-084|Remix/duplicate immutability|Snapshot source record; Remix có destination và cả error/warning branches; duplicate nếu flow được cung cấp.|Success tạo id mới + timestamps mới; source byte snapshot không đổi. Errors hoặc cancelled warning không tạo new record. Duplicate không lén Update source.
APP-08|FR-046–049|Clipboard success/failure|Copy effective text bằng mouse/keyboard; stub clipboard rejection/permission denied.|Success đúng plain raw text; failure thông báo thật và text vẫn selectable để copy thủ công. Không báo Copied khi clipboard thất bại.
APP-09|FR-021–025, FR-034–039, FR-099–102|Disclosure + errors|Mở/đóng multi-groups, navigate-return trong active draft, submit invalid field trong closed group.|Default Product identity và Lighting and composition mở; 3 nhóm còn lại đóng; multi-open và session choices được giữ. Error mở affected group, focus first affected field; hidden controls không vào tab order.
APP-10|FR-113–115|Learn exact field|Mỗi field trigger bằng keyboard/mouse; xem compact panel và full article; quay Build.|Name/definition/effect đúng field; nonmodal không che field/focus trap. Same-tab full article tới exact slug + return context, draft không đổi; native Ctrl-click vẫn theo browser.
APP-11|FR-115|Learn dismiss/focus|Mở Learn rồi Escape, outside click hoặc close control ở ba instance.|Panel đóng; focus trở về đúng trigger theo contract; không mất input đang nhập hoặc draft.
APP-12|FR-116–117|Reset to Preset allowed diff|Customize controlled defaults và product identity/name/destination/output; invoke Reset và confirm.|Dialog nêu Preset/count/scope; chỉ controlled defaults reset. Identity, name, destination, selected Preset, generated/edited text không đổi; không generate/save side effect.
APP-13|FR-116–117|Reset cancel/version|Thử unchanged Preset, cancel/Escape, missing historical catalogue snapshot và Start Blank.|Reset chỉ available khi có controlled diff; cancel giữ nguyên state. Missing version vô hiệu hóa reset có giải thích, không áp defaults version mới hoặc reset Blank như Preset.
APP-14|FR-107–109, FR-118|Selected visual context|Đổi Preset trước Generate, navigate và chọn Project; chuyển missing asset.|Preview/name/intent cập nhật ngay đúng active ID; icon + label cùng selected color, có rail/tint/weight; Project folder fill ngoại lệ được giữ. Fallback không cản input/action.
''')
group('LocalStorageStore và schema migration','Contract + integration',r'''
STO-01|FR-050–051, FR-060, FR-063, FR-079|Round-trip envelope|Write/read records chứa config, generated/edited prompt, stable IDs, versions, timestamps và projectId null/non-null.|Schema version và mọi valid field được bảo toàn qua reload; catalogue builtin không serialize vào user records; async ProjectStore/SavedConfigurationStore contract nhất quán.
STO-02|FR-064, FR-079|Corrupt JSON/envelope|Seed malformed JSON, wrong root shape, bad schema/ID/type/date và duplicate ID; khởi động app.|Guard reject có message; app shell/catalogue vẫn dùng được; không silent clear/reset corrupted bytes. Hành vi xử lý từng bad record vs toàn envelope theo supported storage contract, có evidence.
STO-03|FR-063–064, FR-067|Unsupported version|Seed newer envelope và compiler/profile versions unknown; thử load/regenerate.|Không downgrade hoặc rewrite tự động. Có unsupported path trung thực; bản stored text được bảo toàn, chỉ supported migration được phép thay đổi dữ liệu.
STO-04|FR-063, FR-067|Supported migration|Seed representative prior supported schema fixture; migrate rồi read/update; inject migration exception.|Success đúng reviewed mapping và versions, IDs/text/provenance còn nguyên. Failure giữ nguyên pre-migration envelope; không xóa để app chạy. Những schema chưa có implementation: Blocked, không giả Pass.
STO-05|FR-078–079|Quota/storage unavailable|Inject QuotaExceededError và SecurityError vào create/Update/Remix/migration write; dùng populated old envelope.|Old stored bytes + editor text/config vẫn còn; không Saved hoặc partial destructive result; Copy còn dùng được. Revalidate/fix/retry explicit, không infinite automatic retry.
STO-06|FR-068, FR-078–079|Validate next envelope|Corrupt unrelated record hoặc membership trong next envelope trước write; thử mutate metadata.|Full write payload bị kiểm tra và invalid write bị chặn; stale UI success không che store failure. Error nêu hành động thất bại, không xóa healthy data để chữa lỗi.
STO-07|FR-062, FR-066|Storage seam isolation|Chạy cùng contract tests với in-memory test double và LocalStorageStore; inspect imports.|UI/compiler/validator không import vendor storage internals; Promise rejection được handle. Test double không thay requirement thành database production đã triển khai.
STO-08|FR-060, FR-078|Origin/reload/offline|Reload cùng origin; mở preview/production origins khác; tắt network sau khi app/assets đã load.|Cùng origin đọc dữ liệu persisted; khác origin không được gọi là mất dữ liệu hay automatic sync. Core đã load chạy offline; không hứa cold offline/PWA/service worker mới.
''')
group('Workspace, Projects và My Prompts','Integration + end-to-end',r'''
PRJ-01|FR-080, FR-081–083|Permanent Workspace/null|Khởi tạo empty browser context, create/rename Project, save No Project; inspect records/UI.|Có một Workspace không delete; Project tùy chọn; No Project là projectId:null, không fake entity. Project CRUD không tạo account/workspace selector mới.
PRJ-02|FR-081, FR-088|Project CRUD/filter|Create/rename/delete empty Project, chọn filter và reload; thử blank name theo guard.|CRUD functional, names hợp lệ; filter/membership đúng record. Không đưa Preset/Workflow CRUD vào flow.
PRJ-03|FR-085–086|Delete children destructive|Project có 3 prompts; chọn delete prompts, kiểm tra count/target rồi confirm; instance khác cancel.|Chỉ explicit destructive choice + confirmation xóa Project và đúng 3 children. Cancel giữ tất cả; default không phải destructive delete.
PRJ-04|FR-082, FR-085|Move children → No Project|Project có 3 prompts; giữ safe default move No Project rồi confirm.|Project bị xóa; 3 prompts giữ nội dung/IDs và projectId:null; không orphan foreign key hoặc mất text.
PRJ-05|FR-085–087|Move children → Project|Có source A, targets B/C; chọn B rồi confirm; thử chỉ còn A và target biến mất trước commit.|A bị exclude; khi không có target option disabled. Success chỉ chuyển A children sang B; stale/invalid target không commit partial deletion và báo lỗi.
PRJ-06|FR-055, FR-085–087, FR-078|Delete failure/cancel|Inject write rejection ở từng strategy; thử Escape/cancel Prompt delete và Project delete.|Không fake success; store state cũ nguyên vẹn. UI phản ánh failure; retry không double-delete. Prompt delete đúng item và cancellation giữ item.
LIB-01|FR-056–059, FR-088|Combined search/date/project|Seed names khác case và ngày khác nhau ở null/A/B; kết hợp name + createdAt range + Project.|Name case-insensitive; filters kết hợp AND trên createdAt và membership đúng; empty results rõ ràng. Boundary/timezone được ghi cụ thể theo query contract đã duyệt, không tự phát minh date semantics.
LIB-02|FR-051–059, FR-088|List/empty/no results|Empty store rồi populated filtered store; mở đúng record, đổi filter, delete và reload.|Có empty/no-results states khác nhau, table/action đúng item, search không mutate records; library refresh sau successful write.
''')
group('UX, browser, accessibility và chất lượng giao diện','Manual + end-to-end',r'''
UX-01|FR-089–094|Home → discovery → Build|Fresh load; đi từ Workflow/Preset cards đến Build bằng mouse/keyboard.|Home default; purpose/path rõ; available cards có image/name/Workflow/intent/action. Unavailable status trung thực; Apply không hover-only; non-guarantee wording thấy được.
UX-02|FR-095–098, FR-113–115|English routes/knowledge|Duyệt Home, Build, Built-in, My Prompts, Encyclopedia, dialogs/toasts/validation và exact-field route reload.|English UI/knowledge nhất quán; không language switch giả hoặc research copy chưa dịch. Direct route được kiểm tra theo route contract, draft return được giữ trong active session.
UX-03|FR-099–102|Disclosure status summary|Collapse default/customized/error/warning groups; resolve issues và điều hướng trở lại.|Count/state đúng fields thực tế; error/warning status cập nhật; không chỉ dùng color hoặc hidden content để hiểu trạng thái.
UX-04|FR-103, FR-110–111, FR-118|Visual tokens/icons|So Home/Built-in/Build/My Prompts với v1.4 và UI rules; inspect local asset failures/monochrome.|Minimalist Product Utility, local Lucide outline nominal 18px/2px, labels visible; 7px controls/10px panels. Selected icon+label + non-color cue; không substitute emoji/system glyphs.
UX-05|FR-020–025, FR-102–103, FR-115–116|Keyboard/focus sequence|Tab/Shift+Tab qua toàn flow; dùng Enter/Space, Escape; error auto-reveal, warning/reset/delete dialogs.|Order theo task, focus visible; labels/accessible names đúng; dialog focus vào/ra hợp lý, Learn nonmodal không trap. Closed groups không tabbable; error field và trigger focus chính xác.
UX-06|FR-099–102|Responsive/reflow/zoom|Chạy desktop và samples 320/375/767/768/1280px; zoom 200%; dùng long values/errors và open Learn.|Dưới 768px một cột đúng task order; không mất CTA/text, overlap hoặc page horizontal overflow. Samples là cách kiểm tra boundary, không thêm breakpoint product ngoài baseline.
UX-07|FR-103, FR-118|Contrast và non-color cues|Đo text, focus, selected, error, disabled, panel boundaries trên actual UI.|Ghi ratios và exceptions: normal text target 4.5:1, large text 3:1; meaningful component/state boundaries target 3:1 theo baseline/W3C. Không tuyên bố full WCAG conformance chỉ vì suite này pass.
UX-08|FR-050–059, FR-078–079|Honest feedback/states|Exercise loading/write failure/empty/no results/success/warnings/errors/missing images và Copy failure.|Message nêu đúng outcome, giữ draft khi failure, success chỉ sau completed action; không spinner/toast giả hoặc dead controls.
UX-09|FR-093, FR-109–110, FR-119|Offline/local provenance|Block remote assets/network sau load; inspect image alt, action labels, disclaimer và provenance.|Local assets/knowledge/core còn dùng hoặc fallback rõ; không external API/image-generation requests. Provenance không biến 6 sample images thành proof cho 10 compiler outputs.
UX-10|FR-001–005, FR-040–059, FR-080–088|Chrome/Edge full flow|Trên stable Chrome và stable Edge, ghi version/OS; chạy Generate/edit/Copy/reload/Update/Remix/search/delete.|Cả hai browsers hoàn tất cùng expected behavior; core actions cập nhật trong page, không yêu cầu full reload. Không unhandled exception hoặc divergence chưa giải thích; Firefox/Safari optional không thay mandatory matrix.
UX-11|FR-105–106, FR-112|Scope honesty|Duyệt mọi visible navigation/action và Built-in; inspect production bundle endpoints.|Không Admin, incomplete catalogue CRUD, auth, image upload/generation, API integration hoặc generic form builder. Deferred features không hiển thị như đã hoạt động.
UX-12|FR-040–049, FR-079|Untrusted input display|Nhập HTML/script-like text, Markdown metacharacters và long names vào fields/name/knowledge-search; load saved counterpart.|User content hiển thị an toàn dưới dạng text; không script execution hoặc markup injection. Plain-text output không cần rich Markdown renderer mới; error không lộ unrelated stored content.
''')
group('Release, rollback và pha Database sau MVP','Build / release / deferred',r'''
REL-01|FR-062, FR-066, FR-104|Reproducible production build|Từ clean checkout dùng pinned lockfile; chạy catalogue checks, type check, unit/contract tests và production build.|Không missing references/type/build errors; lưu commit/build ID + dependency versions + test evidence. Mockup HTML không được ghi nhận là production build.
REL-02|FR-001–005, FR-040–059, FR-078, FR-089|Release smoke|Serve built artifact trên intended origin; Home → valid Preset → Generate → edit/Copy → reload → Update → Remix → filter.|Assets/routes/output/persistence đúng, no console failure; release build được chạy chứ không chỉ dev server. Smoke trên Chrome và Edge và mobile-width sample.
REL-03|FR-063, FR-067, FR-078|Rollback/schema drill|Trên isolated test origin, chạy previous build + old data, candidate + migration, rồi previous build với candidate data.|Ghi compatibility thật; rollback code không được giả định rollback storage. Nếu old build không đọc schema mới, dừng promotion hoặc có reviewed recovery path bảo toàn dữ liệu trước release.
REL-04|FR-061, FR-065–067|Database transition gate|Khi DR-005 được giải quyết, chọn vendor/access/ownership/migration; chạy cùng store contracts với DatabaseStore và representative fixtures.|Deferred ở MVP. Trước final submission phải có plan + implementation + evidence đúng DOC-02; không coi LocalStorageStore hoặc this plan là hoàn tất database requirement.
REL-05|FR-071–075, FR-098, FR-106|Future scope gate|Trước thêm Workflow/localization/admin source, đối chiếu research + updated requirements/catalogue và test scope.|Deferred release; có approved contracts và regression trước triển khai. Không gán Pass/Fail MVP cho functionality chưa nằm release Product.
REL-06|FR-069|Relational projection review|Đối chiếu DOC-03 reference relational projection với logical Saved Configuration record và Project membership.|Projection chứa fields/IDs/relations tương ứng logical localStorage record và không giả định vendor. Đây là document/design review; không là bằng chứng database runtime đã được triển khai.
''')

DOC=Document()
sec=DOC.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=Cm(1.8);sec.left_margin=sec.right_margin=Cm(2)
normal=DOC.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10.5)
normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.08
for n,size in [('Title',25),('Subtitle',14),('Heading 1',17),('Heading 2',13),('Heading 3',11)]:
    s=DOC.styles[n];s.font.name='Arial';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
    s.font.italic=False
    s.paragraph_format.keep_with_next=True
    s.paragraph_format.space_before=Pt(12);s.paragraph_format.space_after=Pt(7)
for style in DOC.styles:
    for border in style.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
DOC.core_properties.title='VisualPrompt Studio — DOC-06 Test, Release and Maintenance Plan'
DOC.core_properties.author='VisualPrompt Studio'
DOC.core_properties.subject='Product MVP verification and release governance'
foot=sec.footer.paragraphs[0];foot.alignment=2
foot.add_run('VisualPrompt Studio • DOC-06 v1.0 • ')
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
for r in foot.runs:r.font.size=Pt(8)
def p(text='',bold=False,style=None):
    x=DOC.add_paragraph(style=style);r=x.add_run(text);r.bold=bold;return x
def h(text,level=1):
    x=DOC.add_heading(text,level=level)
    if level==1:x.paragraph_format.page_break_before=text in ['Cách đọc và điều hướng','Phạm vi, trạng thái và nguồn chuẩn','Danh mục test có thể thực thi','Traceability và coverage gate','Kế hoạch release và rollback','Tài liệu tham chiếu']
    return x
def table(headers,rows,widths=None):
    t=DOC.add_table(rows=1, cols=len(headers));t.autofit=False
    widths=widths or [17/len(headers)]*len(headers)
    for i,width in enumerate(widths):t.columns[i].width=Cm(width)
    for i,v in enumerate(headers):t.rows[0].cells[i].text=v
    for row in rows:
        for c,v in zip(t.add_row().cells,row):c.text=str(v)
    for ri,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
        if ri==0:
            hd=OxmlElement('w:tblHeader');trPr.append(hd)
        for i,c in enumerate(row.cells):
            c.width=Cm(widths[i]);pr=c._tc.get_or_add_tcPr()
            shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E7EDF4' if ri==0 else ('F7F8FA' if ri%2==0 else 'FFFFFF'));pr.append(shading)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(4);para.paragraph_format.space_before=Pt(3)
                for r in para.runs:r.font.size=Pt(9);r.bold=ri==0
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders);p();return t
def bullets(items):
    for x in items:p(x,style='List Bullet')
def source_path(text):
    if text.startswith('04_Frontend_Design/Mockups/README.md;'):
        text=text.replace('/README.md;','/\nREADME.md;')
    elif '/' in text:
        folder,filename=text.rsplit('/',1);text=folder+'/\n'+filename
    x=p(text)
    for r in x.runs:r.font.size=Pt(9.5)
def external_link(label,url):
    para=DOC.add_paragraph();para.paragraph_format.keep_with_next=True
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),DOC.part.relate_to(url,RT.HYPERLINK,is_external=True))
    run=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'174A72');pr.append(color)
    run.append(pr);tx=OxmlElement('w:t');tx.text=label;run.append(tx);link.append(run);para._p.append(link)
def refs_expand(s):
    out=set()
    for a,b in re.findall(r'FR-(\d{3})[–-](\d{3})',s):out.update('FR-%03d'%i for i in range(int(a),int(b)+1))
    out.update(re.findall(r'FR-\d{3}',s));return out

p('VisualPrompt Studio',style='Title')
p('DOC-06',style='Subtitle')
p('Test, Release and Maintenance Plan',style='Subtitle')
p('Kế hoạch kiểm thử, phát hành và bảo trì',bold=True)
p('Phiên bản 1.0 • 11/10/2026')
table(['Thuộc tính','Giá trị'],[
 ('Phạm vi','Product Advertising and Photography — browser MVP'),
 ('Trạng thái tài liệu','Hoàn thiện nội dung để review; release policy đề xuất, chưa phải release approval'),
 ('Trạng thái ứng dụng','Chưa có production implementation/build trong baseline đã kiểm tra; application cases: Not Run'),
 ('Vai trò','Chủ dự án review scope/exit gates; người triển khai chạy tests và lưu evidence; người review kiểm tra kết quả'),
 ('Dependency','DOC-01 v2.3; DOC-02 v2.2; DOC-03 v2.1; DOC-04 v1.9; DOC-05 v1.5; Stage 6 v1.4; UI Rules v1.2'),
 ('Version history','v1.0: test matrix, traceability, release/rollback gates, maintenance và Database transition.')],[4.1,12.9])
p('Mục tiêu',bold=True)
p('Chuyển baseline Product đã chốt thành các kiểm tra có điều kiện đầu vào, hành động và kết quả quan sát được. Tài liệu xác định bằng chứng cần có trước khi demo/release và cách bảo toàn cấu hình, Prompt đã lưu khi lỗi storage hoặc thay đổi phiên bản.')
p('Giới hạn bằng chứng',bold=True)
p('372 fixtures trong DISC-01 là kết quả đánh giá nghiên cứu; productionAppTested:false và creativeEffectiveness: Not Empirically Tested. Không dùng chúng để tuyên bố MVP đã pass hoặc hình ảnh sinh ra sẽ giống preview.')

h('Cách đọc và điều hướng')
p('Các Heading có thể dùng trong Navigation Pane của Word/LibreOffice. Mục lục dưới đây liên kết đến các phần chính; danh mục test dùng stable case IDs. Test instances mang parameter và fixtureId riêng để không đánh đồng một case với một lần thực thi.')
TOC_NAMES=['Phạm vi, trạng thái và nguồn chuẩn','Chiến lược và môi trường kiểm thử','Dữ liệu kiểm thử và oracle','Danh mục test có thể thực thi','Traceability và coverage gate','Quản lý kết quả và defects','Kế hoạch release và rollback','Bảo trì và thay đổi phiên bản','Open questions và quyết định còn lại','Tài liệu tham chiếu']
for i,name in enumerate(TOC_NAMES):
    para=DOC.add_paragraph();link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),f'section_{i}')
    run=OxmlElement('w:r');pr=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),'174A72');pr.append(co);run.append(pr)
    tx=OxmlElement('w:t');tx.text=name;run.append(tx);link.append(run);para._p.append(link)
def section(name):
    x=h(name);i=TOC_NAMES.index(name)
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(100+i));start.set(qn('w:name'),f'section_{i}');x._p.insert(0,start)
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(100+i));x._p.append(end)

section(TOC_NAMES[0])
h('MVP được kiểm thử',2)
bullets(['Home, Build, Built-in, My Prompts và Field Encyclopedia; English UI và kiến thức runtime.',
 '30 Product fields; tối đa 10 built-in Presets, Start Blank ngoài count; 10 immutable template profiles và compiler contract product-molds-1.0.0.',
 'Pure validator/compiler; V01–V06 structural errors, V07–V14 art-direction warnings; explicit confirmation và fresh validation trước mutation.',
 'Generate/Remix compile + autosave; load/Update, manual edit, Copy, search/date/Project filters; một Workspace, Projects tùy chọn và No Project = projectId:null.',
 'Versioned LocalStorageStore qua Promise-based store interfaces; corruption/unsupported version/quota paths; keyboard, responsive, Chrome/Edge và latest Learn/reset behavior.'])
h('Ngoài Product MVP và các gate sau',2)
table(['Phạm vi','Disposition'],[
 ('FR-071–075: các Workflow còn lại','Deferred; nghiên cứu fields/rules trước implementation. Không đòi full DOC-02 Workflow suite pass trong Product release.'),
 ('FR-061/065/067: Database phase','FR-061/065 thuộc final submission; FR-067 migration plan vẫn Proposed trong DOC-02. REL-04 là gate pha sau, không claim đã hoàn tất.'),
 ('FR-098 localization; FR-106 future management source','Deferred; không thêm translation dependency, Admin CRUD hoặc generic form builder vào MVP.'),
 ('AI/API/image upload/image generation/auth','Không thuộc core; preview không là generation integration hoặc output-quality benchmark.'),
 ('Vercel hosting','Optional; checklist hosting chỉ áp dụng khi chọn deploy. Tài liệu này không cấp quyền publish.')],[6,11])
h('Authority, source inconsistencies và trạng thái',2)
p('Confirmed: phạm vi Product và các quyết định chủ dự án đã được đưa vào baseline mới. Inferred: test scenarios được suy ra từ invariants, acceptance criteria và failure modes. Assumptions: người triển khai có thể inject lỗi và dùng isolated browser context để kiểm thử. Proposed: test tooling cụ thể, defect policy và exit gates dưới đây cần review; không nâng một Proposed requirement thành Confirmed chỉ vì có test case.')
p('DOC-02 và DOC-05 cùng UI Rules v1.2 là nguồn behavior; DOC-03/04 quyết định state/data/module contracts. Stage 6 v1.4 là visual reference. M03–M06 v1.0 chỉ giữ interaction intent; shell cũ không có authority. Later owner decisions được ghi trong current docs supersede research scopeHint/approved flags cũ; catalogue archival giữ nguyên.')
table(['Vấn đề','Cách trace và kiểm tra'],[
 ('INV-017 trùng trong DOC-03','Profile identity: FR-019/AC-009; regeneration confirmation: FR-077/AC-028/AC-029. Test log phải ghi qualified meaning, không chỉ INV-017.'),
 ('9 research groups vs 5 Build categories','Dùng latest DOC-05/FR-099–102 và v1.4. CAT-03 kiểm tra 30 fields map một lần vào 5 visible categories.'),
 ('Research catalogue chưa runtime-ready','CAT-07 yêu cầu checklist adapter/reconciliation. JSON runtimeReady:false không được tự đổi thành readiness evidence.'),
 ('6 previews vs 10 Presets','CAT-06 yêu cầu stable-ID map hoặc fallback cho tất cả 10 Presets; không invent name equivalence.'),
 ('Proposed architecture/migration requirements','Giữ status gốc trong traceability. Review/acceptance độc lập trước khi dùng làm mandatory final gate.')],[6,11])

section(TOC_NAMES[1])
table(['Layer','Mục tiêu và cách dùng evidence'],[
 ('Unit / build-time','Catalogue integrity, rules, compiler normalization/escaping/determinism, reset diff và filtering; không DOM/storage thật.'),
 ('Contract / integration','Store envelope, guards, async error paths, migration, membership và mutation sequencing. Inject write failures bằng test double có cùng interface.'),
 ('End-to-end','Build → validation → confirm → compile → persist → reload; Project deletion; Learn/reset/focus. Kiểm tra bằng UI trên built application.'),
 ('Manual / visual','English content, selected cues, contrast, narrow/zoom, preview semantics và keyboard usability. Evidence gồm screenshots + notes, không chỉ mockup comparison.'),
 ('Release smoke / rollback','Built artifact trên intended origin, schema compatibility và rollback rehearsal. Không thay thế regression suite.')],[4.5,12.5])
p('Tooling đề xuất tối thiểu: runner tương thích TypeScript cho pure units, browser automation cho các core flows và manual check cho visual/focus/contrast. Chọn runner khi scaffold Vite/TypeScript được tạo; không thêm framework/reducer/router/service layer chỉ để thuận tiện testing. Cùng AppState, named update functions và store seams đã chốt tiếp tục là architecture baseline.')
table(['Môi trường','Điều kiện và dữ liệu phải ghi'],[
 ('Mandatory browsers','Stable Chrome và stable Edge trên môi trường Windows thực thi; ghi exact version, OS, build ID và origin. Không chốt số version chưa được chạy.'),
 ('Responsive samples','320, 375, 767, 768, 1280 px và 200% zoom; samples kiểm tra behavior/boundary, không thành breakpoint specification mới.'),
 ('Storage contexts','Clean isolated test origin; populated valid origin; injected quota/security error; corrupt/unsupported/migration fixtures. Không clear LocalStorage thật của chủ dự án.'),
 ('Network/assets','Normal load; missing local image; network off sau khi assets đã load. Cold offline không thuộc accepted promise.'),
 ('Optional','Firefox/Safari, Vercel preview/production và device bổ sung nếu được chọn; không che kết quả fail của Chrome/Edge.')],[4.5,12.5])
p('Không đặt performance SLA hoặc dung lượng dataset tùy ý thành requirement. Khi chạy, ghi representative dataset size, startup/search/write timing và perceived delay; nếu cần target mới, đưa về requirements review trước khi dùng làm release blocker.')

section(TOC_NAMES[2])
table(['Data pack','Input cụ thể và oracle'],[
 ('D-01 valid Product base','Required subject/use-case/name + explicit neutral profile hoặc approved Preset defaults. Không chứa real user data; lưu input JSON và expected Markdown độc lập.'),
 ('D-02 catalogue/profile set','30 fields, 10 Presets, 10 profiles, V01–V14; curated matching/non-matching operand pairs. Ghi version/hash và runtime adapter version.'),
 ('D-03 text boundaries','Fixed/suggested/free values, max−1/max/max+1 theo limit của mỗi field; Unicode scalar, NFC/decomposed, controls, braces, dollar, quotes, backslash và whitespace.'),
 ('D-04 saved record set','No Project/A/B/C; distinct IDs/names/createdAt/updatedAt, generated-only/manual-edited, provenance và saved versions; old-supported/newer/corrupt envelope variants.'),
 ('D-05 failures','Injected quota/security rejection, migration exception, invalid next-envelope, target Project mất trước write, Copy rejection và missing asset.'),
 ('D-06 research fixture ports','372 research fixture IDs. Mỗi port có concrete input + expected issues/output + source rule/profile + proof assertion; missing input không được đoán rồi tái dùng research Pass.')],[4.5,12.5])
p('Pairwise/rule-oriented coverage tránh Cartesian product không có lý do. Chạy từng field boundary, từng Preset/profile, mỗi rule match và non-match; thêm errors+warnings cùng lúc, stale confirmation và write failure. Những counts là số test instances sẽ xác định khi port fixtures, không phải số 372 production tests mặc định.')
p('Oracle ưu tiên: approved compiler/rule contract → acceptance criteria/invariants qualified → independent fixture expected values. Không lấy actual từ implementation để làm expected của chính implementation. Byte assertions dùng cho deterministic Markdown/exact-copy; state diff dùng cho reset/update/remix/cancel; stored-byte snapshot dùng cho rejected writes.')
p('Date filtering: dùng createdAt theo baseline, case-insensitive name và AND composition. Định nghĩa boundary/timezone chưa đủ cụ thể phải được ghi trong query contract trước executable date assertions; không biến một timezone assumption thành product rule.')

section(TOC_NAMES[3])
p('Mọi case dưới đây đang Not Run. Action mô tả setup + thao tác; expected là observable oracle. Với case parameterized, tạo một execution row cho mỗi parameter/fixture và mỗi mandatory browser khi có UI. Test level ở từng nhóm xác định evidence cần lưu.')
for name,level,cases in GROUPS:
    h(name,2);p('Layer: '+level+' • Trạng thái khởi tạo: Not Run')
    for c in cases:
        x=p(c['id']+' — '+c['title'],bold=True);x.paragraph_format.keep_with_next=True
        x=p('Trace: '+c['refs']+' | Action: '+c['action']);x.paragraph_format.keep_with_next=True
        p('Expected: '+c['expected'])

section(TOC_NAMES[4])
p('Bảng dưới map từng FR thật trong DOC-02 sang case IDs. Mapping không đồng nghĩa coverage đã thực thi; trước release, mỗi in-scope FR phải có actual + evidence và disposition rõ. Acceptance criteria được dùng làm oracle đi kèm FR; bổ sung AC ID cụ thể vào execution log, đặc biệt AC-009, AC-028/029 và AC-069–075 cho latest behavior.')
src=Document(ROOT/SOURCES[1][1]);reqs={}
for t in src.tables:
    for row in t.rows:
        cells=[c.text for c in row.cells]
        if re.fullmatch(r'FR-\d{3}',cells[0].strip()):reqs[cells[0].strip()]=cells
allcases=[c for _,_,cs in GROUPS for c in cs]
coverage={r:[c['id'] for c in allcases if r in refs_expand(c['refs'])] for r in sorted(reqs)}
missing=[r for r,v in coverage.items() if not v]
unknown=sorted(set().union(*(refs_expand(c['refs']) for c in allcases))-set(reqs))
# Numerical ranges deliberately omit IDs unused by the source; never invent requirements.
unknown_referenced=[r for r in unknown if any(r in re.findall(r'FR-\d{3}',c['refs']) for c in allcases)]
assert not missing,missing
assert not unknown_referenced,unknown_referenced
deferred={'FR-061','FR-065','FR-067','FR-071','FR-072','FR-073','FR-074','FR-075','FR-098','FR-106'}
rows=[]
for r,caseids in coverage.items():
    original=' / '.join(reqs[r][2:])
    status='Deferred/gate pha sau' if r in deferred else ('Proposed — giữ status gốc' if 'Proposed' in original else 'Product release')
    rows.append((r,', '.join(caseids),status))
table(['FR ID','Cases kế hoạch','Disposition'],rows,[2.3,10.3,4.4])
h('Coverage các nguồn khác',2)
table(['Nguồn','Đối chiếu bắt buộc'],[
 ('DOC-03 invariants','INV-001–016: CAT/APP/STO/PRJ; INV-017 phải qualified theo profile vs regen; INV-018–020: PRJ/VAL. Không renumber source trong test plan.'),
 ('DOC-04 FE/ADRs','Named state transitions và pure modules: CMP/APP; store seams: STO; Learn/reset decisions: APP-10–13. REL-01 inspect dependencies/boundaries.'),
 ('DOC-05 UXAC','UXAC-001–012: core UX/validation/persistence; UXAC-021–040: Home/preview, 5 groups, Learn/reset/selection/provenance. APP-09–14 và UX-01–12 là scenario groups; execution log nêu exact criterion.'),
 ('NFR/BRULE','Native TS boundaries/purity/offline: CMP/STO/REL; keyboard/Chrome/Edge: UX-05/10; API/auth/deployment scope: UX-11/REL. BRULE record/source/gates được assert qua state/store snapshots.'),
 ('Excluded / Proposed','Không mark deferred Workflow/DB/localization/future Admin như missing MVP. Proposed requirements vẫn yêu cầu owner review; no fabricated approval.')],[4.5,12.5])
p(f'Tổng catalogue kế hoạch: {len(allcases)} stable case IDs trong {len(GROUPS)} nhóm; {len(reqs)} FR hiện diện trong source được map. Đây là planned coverage theo FR, chưa phải executed coverage hoặc pass percentage. Expanded test instances sẽ nhiều hơn khi port inputs và chạy browsers.')
h('Non-functional requirements',2)
NFR_MAP={
 'NFR-001':('REL-01','Inspect production dependencies/source: vanilla HTML/CSS/TS; không framework tương đương.'),
 'NFR-002':('REL-01, UX-04','Review runtime dependencies; browser/native TypeScript preference, local asset bundle có scope rõ.'),
 'NFR-003':('CMP-02, UX-09, UX-11','Configure/validate/compile không cần AI hoặc image-provider network call.'),
 'NFR-004':('STO-07, REL-01','Inspect models/data/validation/compiler/storage/UI responsibility boundaries.'),
 'NFR-005':('UX-10, APP-09','Core actions cập nhật trong page; manual reload chỉ là persistence test.'),
 'NFR-006':('UX-05, APP-10–13','Semantic controls, mouse/keyboard, logical tab order/focus; không full WCAG claim.'),
 'NFR-007':('UX-10, REL-02','Stable Chrome/Edge execution evidence; ghi versions thực tế.'),
 'NFR-008':('PRJ-01, UX-11','Core không account/profile-data prerequisite.'),
 'NFR-009':('CMP-02, VAL-01–17','Pure deterministic units; validator issues repeatable cho cùng input, không cần DOM.'),
 'NFR-010':('REL-02, REL-03','Optional hosting; Vercel preferred nếu chọn deploy; no implied publish approval.'),
 'NFR-011':('STO-07, REL-01','Review code structure theo architecture cơ bản đã chốt, không enterprise patterns mới.'),
 'NFR-012':('UX-09, UX-11, REL-05','External API deferred sau MVP hoàn tất + kiểm thử; không dependency của release này.')}
source_nfr={r.cells[0].text.strip() for t in src.tables for r in t.rows if re.fullmatch(r'NFR-\d{3}',r.cells[0].text.strip())}
assert set(NFR_MAP)==source_nfr
table(['NFR ID','Cases','Oracle bổ sung'],[(k,*v) for k,v in NFR_MAP.items()],[2.3,4.2,10.5])

section(TOC_NAMES[5])
h('Execution record và trạng thái',2)
table(['Trường log','Nội dung bắt buộc'],[
 ('Identity','runId, caseId, instance/parameter ID, research fixtureId nếu có, exact FR/AC/invariant qualified, test layer.'),
 ('Environment','commit/build ID, catalogue/compiler/schema versions, OS/browser version, origin, viewport/zoom, dataset ID/hash.'),
 ('Procedure/result','preconditions + input, action, expected, actual, Pass/Fail/Blocked/Not Run, execution time và tester/reviewer role.'),
 ('Evidence','Automated assertion/log hoặc screenshot/state/store diff có link; defect ID và retest result. Không log entire real-user prompt contents.'),
 ('Release disposition','In-scope/deferred/proposed, known defect/waiver owner và lý do; evidence thiếu không được chuyển Pass.')],[4.5,12.5])
p('Not Run: chưa chạy implementation. Blocked: thiếu build/input/contract hoặc environment làm không chạy được. Pass: actual đáp ứng expected với evidence ở đúng build. Fail: observable mismatch. Một case chỉ tổng hợp Pass khi tất cả mandatory instances pass; optional environment kết quả giữ riêng.')
h('Defect policy đề xuất',2)
table(['Mức','Ví dụ và xử lý release'],[
 ('Critical','Data loss, wrong-project destructive delete, executable input injection, corrupt migration. Block release; preserve evidence/bytes, fix + targeted/full relevant regression.'),
 ('High','Error bypass, stale warning confirmation, wrong compiler/profile, failure false success, broken main flow hoặc mandatory browser. Block release.'),
 ('Medium','Usable but incorrect secondary behavior, field help mismatch, keyboard issue có workaround. Review impact; mandatory requirement chưa đạt vẫn phải fix hoặc documented owner waiver.'),
 ('Low','Non-blocking copy/spacing discrepancy không phá function/access. Có thể defer với explicit issue/owner; không tự hạ severity của accessibility/data defect.')],[4,13])
p('Không đặt response SLA hoặc người phụ trách chưa được chỉ định. Chủ dự án quyết định waiver đối với policy đề xuất; waiver không đổi Failed thành Passed. Retest liên kết defect với case/instances và build mới, giữ lịch sử actual cũ.')

section(TOC_NAMES[6])
h('Entry và exit gates đề xuất',2)
table(['Gate','Evidence và quyết định'],[
 ('G-01 Entry','Production scaffold/source chạy được; runtime catalogue reconciled + build-time guards; inputs/oracles cho cases mandatory; frozen scope/version references.'),
 ('G-02 Functional','Mọi in-scope mandatory FR/AC có executed evidence trên candidate build; validator/compiler/store tests pass; Chrome/Edge core paths pass; no Critical/High defects.'),
 ('G-03 Data safety','Quota/corrupt/unsupported writes, three Project deletion strategies, Update/remix identity và supported migrations verified; rollback/schema drill documented.'),
 ('G-04 UX','Keyboard/focus, <768 behavior/200% zoom, English UI/knowledge, Learn/reset/selection, local preview fallback/provenance checked against latest baseline.'),
 ('G-05 Review','Chủ dự án review remaining defects/waivers và release scope; no false readiness claims. DOC-06 complete hoặc research fixture pass không tự satisfy gate này.'),
 ('G-06 Final database phase','DR-005 resolved và REL-04 evidence trước final submission; không chặn browser Product MVP chỉ vì DatabaseStore chưa chọn.')],[3.7,13.3])
h('Runbook từ candidate tới demo/release',2)
for text in [
 '1. Freeze candidate: lưu commit/build ID, lockfile, catalogue/profile/compiler/schema versions và change list. Giữ prior known-good build và compatibility notes.',
 '2. Clean build: chạy integrity/type/unit/store checks và production build; lưu outputs. Không dùng Stage 6 mockup files làm artifact ứng dụng.',
 '3. Isolated test: seed synthetic fixtures, chạy mandatory instances, Chrome/Edge, responsive/manual checks. Không clear browser data thật hoặc dùng production prompts làm test dataset.',
 '4. Review evidence: mọi failed/blocked/not-run mandatory instances có disposition; Critical/High unresolved chặn release. Record owner decision cho Proposed gates/waivers.',
 '5. Intended origin smoke: serve đúng built artifact và chạy REL-02; kiểm tra direct routes/assets/local storage/Clipboard behavior. Nếu hosting optional được chọn, preview origin không chứng minh dữ liệu đã chuyển production.',
 '6. Promotion chỉ sau authorization và gate review: ghi release ID/date, scope/known limitations, test summary và rollback target. Không publish chỉ vì document này có checklist.',
 '7. Post-release smoke: trên production origin kiểm tra Home/Build/Copy/storage và thiếu assets; nếu regression material, giữ evidence và dùng rollback/recovery procedure. Không thêm telemetry hoặc external API vào MVP để theo dõi.']:p(text)
h('Rollback và phục hồi dữ liệu',2)
p('Trigger đề xuất: data mutation sai, blocking main flow, error bypass hoặc mandatory-browser regression sau promotion. Trước rollback, xác định affected build/schema/origin, chụp test evidence và bảo toàn envelope qua công cụ kiểm thử/manual backup được kiểm soát. Đây không phải requirement thêm UI export/import feature.')
p('Code rollback chỉ đổi application artifact. LocalStorage của browser/origin không tự quay lại phiên bản dữ liệu trước đó; schema rollback compatibility phải được drill ở REL-03. Nếu old build không hỗ trợ new schema, không ép load bằng clear storage: chọn reviewed compatible recovery build/migration và bảo toàn bytes. Không hứa phục hồi dữ liệu đã bị người dùng/browser xóa khi không có backup.')
p('Nếu chọn Vercel: xác định eligible previous production deployment theo plan/account và current provider documentation; review deployment/domains/config trước Instant Rollback. Rollback phục hồi previous build, không tự áp environment changes. Sau đó chạy smoke và kiểm tra cơ chế promotion tiếp theo. Chi tiết thao tác/provider eligibility phải được kiểm tra lại lúc release [E-03].')

section(TOC_NAMES[7])
table(['Loại thay đổi','Bắt buộc trước merge/release'],[
 ('Bugfix UI/help','Targeted APP/UX regression + core smoke; field help exact ID, English copy, no state loss; update only affected docs/source.'),
 ('Rule/defaults/profile','Stable IDs + reviewed version; CAT/VAL/CMP regression, all affected fixture ports và reset snapshot behavior. Không rewrite immutable historical profile để regenerated saved outputs đổi ngầm.'),
 ('Compiler','New supported compilerVersion + deterministic oracle review, escaping/normalization/exact-copy regression; saved generated/edited text giữ nguyên trừ explicit reviewed migration.'),
 ('Schema/store','Versioned envelope + representative old/new/invalid fixtures, rejected-write preservation, migration/rollback drill. Unsupported future versions không bị silent reset.'),
 ('Dependencies/build','Minimal justified change; type/build checks, Chrome/Edge smoke, local-assets/offline/failure paths. Không thêm runtime dependency không có scope benefit.'),
 ('Project/library behavior','Re-run all membership/deletion strategies/filter AND + write-failure tests; stable record identity và no orphan/no fake null entity.'),
 ('Scope expansion/Database','Resolve corresponding gates, update requirements/domain/architecture/UX/test matrix theo dependency order trước release. Không reuse Product evidence cho unresearched Workflow.')],[4.5,12.5])
p('Maintenance theo change/release trigger, không invent lịch tuần/tháng hoặc SLA. Issue log giữ expected/actual/build/schema/evidence và owner khi được chỉ định. Recovery ưu tiên bảo toàn người dùng đang nhập và bản đã lưu; không sửa corruption bằng blanket storage clear.')
p('Catalogue archival, runtime adapter, executable fixtures và delivery docs có vai trò riêng nhưng cùng canonical stable-ID contract. Documentation version theo policy hiện hành; không áp dụng in-place filename policy mới chưa được chấp nhận. Mỗi revision ghi source versions thay đổi và superseded decisions; không bump DOC-01–05 nếu không có nội dung sửa.')
p('Future DatabaseStore phải chạy cùng Promise-based ProjectStore/SavedConfigurationStore contract trước UI integration. Vendor/access/ownership/migration là DR-005; plan này không chọn backend, auth model hoặc migration API thay chủ dự án. Migration thử trên synthetic representative fixtures trước dùng dữ liệu thật.')

section(TOC_NAMES[8])
table(['ID / priority','Trạng thái và ảnh hưởng'],[
 ('OQ-001 / Medium','Partially resolved: Product 30 fields đã chốt. Remaining Workflow fields cần research trước release sau; không chặn Product tests.'),
 ('OQ-003 / Medium','Partially resolved: Product 14 rules + 372 research fixtures đã chốt. Port executable inputs/oracles vào implementation; remaining Workflow rules research sau.'),
 ('DR-005 / Deferred MVP; High trước final DB','Chọn vendor/access/ownership/migration trước DatabaseStore/final submission. Giữ store seam hiện tại; không giả định đã resolved.')],[5,12])
p('Không tạo thêm owner Open Question từ catalogue reconciliation hoặc thiếu production build: đây là work item rõ theo baseline. Planned exit gates/tooling/waivers là Proposed review items, không new product decisions được tự chốt. Số open items mới: 0; carried items: 3.')
table(['Work item','Hoàn tất khi'],[
 ('IMP-01 Scaffold và runtime adapter','Production Vite/vanilla TS source/build tồn tại; CAT-02–07 và scope source reconciliation chạy được.'),
 ('IMP-02 Fixture ports/oracles','372 fixture records có mapping; missing inputs được bổ sung có source oracle và review; không transplant research actual.'),
 ('IMP-03 Test execution evidence','Run/instance log + screenshots/state/store assertions ở actual candidate build; case statuses cập nhật có trace.'),
 ('REV-01 Review DOC-06 release policy','Owner chấp nhận/điều chỉnh Proposed gates và defect policy trước first release; no formal approval inferred từ việc tạo DOCX.')],[5,12])

section(TOC_NAMES[9])
for label,relative in SOURCES:
    p(label,bold=True);source_path(relative)
for label,relative in [
 ('UI Rules v1.2','04_Frontend_Design/Style_Decisions/Minimalist/VisualPrompt_Studio_UI_Generation_Rules_v1.2.md'),
 ('Stage 6 authority / review','04_Frontend_Design/Mockups/README.md; VisualPrompt_Studio_Stage6_Mockup_Review_v1.4.md'),
 ('DISC-01 research source','DISC01_Configuration_Catalogue_v1.0.json; DISC01_Preset_Molds_v1.0.md; DISC01_Field_Encyclopedia_v1.0.md'),
 ('Process authority','09_Sources_and_Working_Files/Documentation_Process/docs-writer-v3.md')]:p(label,bold=True);source_path(relative)
p('Nguồn kỹ thuật ngoài project — kiểm tra ngày 11/10/2026',bold=True)
for label,url,note in [
 ('E-01 MDN Web Storage API','https://developer.mozilla.org/en-US/docs/Web/API/Web_Storage_API','LocalStorage gắn với origin; browser state không đồng bộ giữa preview/production origins.'),
 ('E-02 W3C Understanding Contrast Minimum','https://www.w3.org/WAI/WCAG22/Understanding/contrast-minimum.html','Normal text 4.5:1 và large text 3:1; kiểm tra riêng không là chứng nhận full WCAG.'),
 ('E-03 Vercel Instant Rollback','https://vercel.com/docs/instant-rollback','Previous production build rollback; eligibility/config phải verify tại release, browser storage recovery là trách nhiệm ứng dụng.')]:
    external_link(label,url);p(note)
OUT.parent.mkdir(parents=True,exist_ok=True)
DOC.save(OUT)
# Apply the documented project presentation baseline to each freshly built file.
import importlib.util
format_spec=importlib.util.spec_from_file_location('project_document_formatting',WORK.parent/'Formatting_Audit_Work/normalize_formatting.py')
format_module=importlib.util.module_from_spec(format_spec)
format_spec.loader.exec_module(format_module)
format_module.normalize(OUT,'DOC 06','Test, Release and Maintenance Plan')
audit={'document':str(OUT),'caseCount':len(allcases),'caseGroupCount':len(GROUPS),'frCount':len(reqs),'nfrCount':len(NFR_MAP),'unmappedFRs':missing,'unknownExplicitFRs':unknown_referenced,'initialStatus':'Not Run','newOpenItems':0,'carriedOpenItems':['OQ-001','OQ-003','DR-005'],'cases':allcases,'coverage':coverage,'nfrCoverage':NFR_MAP,'sources':[{'label':a,'path':b,'sha256':hashlib.sha256((ROOT/b).read_bytes()).hexdigest()} for a,b in SOURCES]}
(WORK/'doc06_plan_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps({k:v for k,v in audit.items() if k not in ['cases','coverage','nfrCoverage','sources']},ensure_ascii=False))
