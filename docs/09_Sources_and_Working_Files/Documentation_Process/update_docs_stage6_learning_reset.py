from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import sys

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.text import WD_BREAK
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[4]
sys.path.insert(0, str(ROOT / "work"))

from build_doc01 import add_body, add_bullets, add_heading, add_table  # noqa: E402


DOCS = [
    {
        "id": "DOC-01",
        "old": "v2.2",
        "new": "v2.3",
        "src": ROOT / "outputs/VisualPrompt_Studio/00_Project_Management/Project_Overview/VisualPrompt_Studio_DOC-01_Product_and_Project_Brief_v2.2.docx",
        "dst": ROOT / "outputs/VisualPrompt_Studio/00_Project_Management/Project_Overview/VisualPrompt_Studio_DOC-01_Product_and_Project_Brief_v2.3.docx",
        "history": "Accepted contextual field learning, safe Reset to Preset semantics, persistent selected-icon treatment and explicit AI preview-asset provenance.",
    },
    {
        "id": "DOC-02",
        "old": "v2.1",
        "new": "v2.2",
        "src": ROOT / "outputs/VisualPrompt_Studio/01_Requirements/Product_Requirements/VisualPrompt_Studio_DOC-02_Product_and_Functional_Requirements_Specification_v2.1.docx",
        "dst": ROOT / "outputs/VisualPrompt_Studio/01_Requirements/Product_Requirements/VisualPrompt_Studio_DOC-02_Product_and_Functional_Requirements_Specification_v2.2.docx",
        "history": "Added testable exact-field Learn, Reset to Preset and persistent selected-icon requirements plus preview-asset disclosure.",
    },
    {
        "id": "DOC-03",
        "old": "v2.0",
        "new": "v2.1",
        "src": ROOT / "outputs/VisualPrompt_Studio/02_Architecture/Data_Architecture/VisualPrompt_Studio_DOC-03_Domain_State_and_Data_Design_v2.0.docx",
        "dst": ROOT / "outputs/VisualPrompt_Studio/02_Architecture/Data_Architecture/VisualPrompt_Studio_DOC-03_Domain_State_and_Data_Design_v2.1.docx",
        "history": "Defined exact-field knowledge routing and a derived Preset reset baseline without adding a persisted database entity.",
    },
    {
        "id": "DOC-04",
        "old": "v1.8",
        "new": "v1.9",
        "src": ROOT / "outputs/VisualPrompt_Studio/02_Architecture/Frontend_Architecture/VisualPrompt_Studio_DOC-04_Frontend_Architecture_and_ADRs_v1.8.docx",
        "dst": ROOT / "outputs/VisualPrompt_Studio/02_Architecture/Frontend_Architecture/VisualPrompt_Studio_DOC-04_Frontend_Architecture_and_ADRs_v1.9.docx",
        "history": "Added contextual-help routing, pure reset-diff behavior and selected-state token rules with no new runtime dependency.",
    },
    {
        "id": "DOC-05",
        "old": "v1.4",
        "new": "v1.5",
        "src": ROOT / "outputs/VisualPrompt_Studio/04_Frontend_Design/UX_and_Interaction/VisualPrompt_Studio_DOC-05_UX_and_Interaction_Specification_v1.4.docx",
        "dst": ROOT / "outputs/VisualPrompt_Studio/04_Frontend_Design/UX_and_Interaction/VisualPrompt_Studio_DOC-05_UX_and_Interaction_Specification_v1.5.docx",
        "history": "Specified compact Learn popovers, exact-field full articles, safe Reset to Preset confirmation and selected-icon state.",
    },
]


VERSIONS = {
    "DOC 01 ": "v2.3",
    "DOC 02 ": "v2.2",
    "DOC 03 ": "v2.1",
    "DOC 04 ": "v1.9",
    "DOC 05 ": "v1.5",
}


def all_paragraphs(doc: DocumentType):
    for paragraph in doc.paragraphs:
        yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in doc.sections:
        for story in (section.header, section.footer):
            yield from story.paragraphs
            for table in story.tables:
                for row in table.rows:
                    for cell in row.cells:
                        yield from cell.paragraphs


def replace_in_runs(doc: DocumentType, replacements: dict[str, str]) -> None:
    for paragraph in all_paragraphs(doc):
        for run in paragraph.runs:
            for old, new in replacements.items():
                if old in run.text:
                    run.text = run.text.replace(old, new)


def update_version_history(doc: DocumentType, version: str, description: str) -> None:
    for table in doc.tables:
        if not table.rows:
            continue
        header = " ".join(table.cell(0, 0).text.split()).casefold()
        if header not in {"phiên bản", "version"}:
            continue
        table._tbl.append(deepcopy(table.rows[-1]._tr))
        row = table.rows[-1]
        values = [version, "11/10/2026", "Codex + LeVy", description]
        for index, cell in enumerate(row.cells):
            cell.text = values[index] if index < len(values) else ""
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8.5)
        return


def update_related_versions(doc: DocumentType) -> None:
    for table in doc.tables:
        for row in table.rows:
            if len(row.cells) < 2:
                continue
            name = " ".join(row.cells[0].text.split())
            for prefix, version in VERSIONS.items():
                if name.startswith(prefix):
                    row.cells[1].text = version
                    for paragraph in row.cells[1].paragraphs:
                        for run in paragraph.runs:
                            run.font.name = "Arial"
                            run.font.size = Pt(8.5)
                    break


def add_decisions(doc: DocumentType) -> None:
    add_table(
        doc,
        ["ID", "Decision", "Status", "Consequence"],
        [
            ["DR-014", "Build groups configurable fields into validation-aware collapsible categories.", "Accepted", "Users scan by category; structural errors and warnings still reveal the exact correction location."],
            ["DR-015", "Primary navigation and high-value actions use one consistent icon family with visible labels.", "Accepted", "Recognition improves without icon-only ambiguity or decorative proliferation."],
            ["DR-016", "MVP catalogue authoring remains versioned TypeScript or JSON with schema validation. Protected Admin and generic form building are deferred.", "Accepted", "The product ships sooner while stable IDs and repository boundaries preserve a future management path."],
        ],
        widths=[0.72, 2.65, 0.75, 2.35],
    )


def append_doc01(doc: DocumentType) -> None:
    add_heading(doc, "Confirmed Build grouping and catalogue-management baseline", level=1)
    add_body(doc, "This section is authoritative where earlier wording or mockups imply one uninterrupted Build form, text-only navigation or an MVP administration surface.")
    add_decisions(doc)
    add_heading(doc, "MVP scope boundary", level=2)
    add_bullets(doc, [
        "Build divides Workflow fields into collapsible categories. The first one or two groups open by default; later groups show field count and default or customization status.",
        "Structural errors and art-direction warnings both identify their affected group and fields. Errors block until corrected; warnings require confirmation and do not become Preset-conflict violations.",
        "Icons support navigation and high-value actions, always with visible text labels in the MVP.",
        "Workflow and Preset definitions remain source-authored and build-time validated for the first MVP.",
        "Any future Admin area is protected and separate from the user's Workspace. A generic drag-and-drop form builder is not part of the course-project MVP.",
    ])


def append_doc02(doc: DocumentType) -> None:
    add_heading(doc, "Grouped Build and catalogue-authoring requirements", level=1)
    add_body(doc, "The following requirements are normative for implementation and supersede any visual reference that renders configuration as one uninterrupted form.")
    add_table(
        doc,
        ["ID", "Requirement", "Priority", "Verification"],
        [
            ["FR-099", "Build shall render Workflow fields in ordered, labelled collapsible categories defined by the active Workflow catalogue.", "Must", "Every visible field belongs to one category and category order matches catalogue order."],
            ["FR-100", "Each category summary shall expose a compact status including field count and, where relevant, default, customized, warning or error state.", "Must", "Collapsed groups remain understandable without opening them."],
            ["FR-101", "The first one or two task-critical groups shall open by default; user-controlled open state shall persist for the active draft session.", "Must", "Navigation and validation do not reset unrelated disclosure choices."],
            ["FR-102", "Structural validation failure shall open the first affected group and expose every affected field; warning confirmation shall map to the same group and field locations.", "Must", "Users can locate and correct every reported issue."],
            ["FR-103", "Primary navigation and high-value actions shall use one consistent icon family with visible text labels.", "Should", "Icons never become the sole accessible name and retain one visual system."],
            ["FR-104", "MVP Workflow, field, Preset, validation and knowledge catalogues shall be versioned source data with stable IDs and build-time integrity validation.", "Must", "Broken references or duplicate IDs fail validation before release."],
            ["FR-105", "The MVP shall not expose Workflow or Preset administration controls inside Workspace.", "Must", "Ordinary navigation contains no Admin or catalogue-authoring action."],
            ["FR-106", "A future management tool shall write the same canonical catalogue contract consumed by the runtime rather than a second representation.", "Should", "The catalogue repository boundary accepts either validated static data or a later persisted source."],
        ],
        widths=[0.62, 2.95, 0.62, 2.3],
    )
    add_heading(doc, "Additional business rules", level=2)
    add_bullets(doc, [
        "BRULE-034: Category disclosure changes visibility only; it never changes configuration values.",
        "BRULE-035: Collapsed controls are removed from keyboard focus order and reopening a category restores a predictable focus sequence.",
        "BRULE-036: A Preset belongs to one Workflow and may reference only field IDs defined by that Workflow.",
        "BRULE-037: Catalogue publication fails when Workflow, field-group, field, Preset, validation-rule or knowledge-entry references are invalid.",
        "BRULE-038: Admin authentication, authorization, publishing workflow, asset upload and generic form building remain deferred until separately approved.",
    ])
    add_heading(doc, "Additional acceptance criteria", level=2)
    add_bullets(doc, [
        "AC-057: Given a fresh Product Build, when the form appears, then Product identity and Lighting and composition are open and later groups expose summaries while collapsed.",
        "AC-058: Given a structural error in a collapsed group, when Generate is attempted, then that group opens and focus reaches its summary or first invalid field.",
        "AC-059: Given an art-direction warning, when the warning is shown, then affected fields and groups are identified and generation can continue only after confirmation.",
        "AC-060: Given keyboard-only use, when focus reaches a group summary, then Enter or Space toggles it and hidden controls are skipped.",
        "AC-061: Given primary navigation, when labels or icons are reviewed, then every icon has adjacent text and no icon-only interpretation is required.",
        "AC-062: Given the MVP build, when ordinary user navigation is inspected, then no Workflow or Preset Admin surface is available.",
    ])


def append_doc03(doc: DocumentType) -> None:
    add_heading(doc, "Canonical Workflow catalogue contracts", level=1)
    add_body(doc, "The catalogue is read-only runtime input in MVP. It is shaped like a future persisted model, but it is not stored in localStorage and does not create an Admin requirement.")
    add_table(
        doc,
        ["Contract", "Required shape", "MVP source", "Invariant"],
        [
            ["Catalogue envelope", "schemaVersion, catalogueVersion, workflows[]", "Versioned TypeScript or JSON", "Schema and catalogue versions are explicit and validation completes before runtime use."],
            ["Workflow definition", "id, name, status, fieldGroups[], presetIds[], validationRuleIds[]", "Static catalogue", "Stable unique ID; every referenced entity exists in the same Workflow scope."],
            ["Field group", "id, workflowId, label, order, fieldIds[], defaultOpen", "Static catalogue", "Order is deterministic; every field belongs to exactly one group in its Workflow."],
            ["Field definition", "id, workflowId, valueType, controlKind, options, constraints, knowledgeEntryId", "Static catalogue", "Control rendering and validation derive from typed metadata, not page-specific branching."],
            ["Preset definition", "id, workflowId, name, defaultsByFieldId, templateProfileId, previewAssetId", "Static catalogue", "Defaults reference only fields owned by the same Workflow."],
            ["Catalogue issue", "code, severity, entityId, path, message", "Build-time validation output", "Errors block catalogue publication; warnings remain visible review evidence."],
        ],
        widths=[1.08, 2.3, 1.05, 2.0],
    )
    add_heading(doc, "Authoring and future management", level=2)
    add_bullets(doc, [
        "MVP catalogue definitions are imported through CatalogueRepository and treated as immutable after initialization.",
        "localStorage stores user Workspace aggregates only; it does not store built-in Workflow or Preset definitions.",
        "A future Admin application may persist CatalogueEnvelope in a database, but must pass the same schema and referential validators before publication.",
        "Field groups provide the configuration-disclosure structure. Runtime open or closed state remains session UI state and is not catalogue ownership data.",
        "Deleting or changing a future Workflow, field or Preset requires impact validation against references and published catalogue versions.",
    ])


def append_doc04(doc: DocumentType) -> None:
    add_heading(doc, "Grouped configuration and catalogue seams", level=1)
    add_table(
        doc,
        ["Module or component", "Responsibility", "Dependencies", "Constraint"],
        [
            ["Configuration group", "Render a semantic category summary, status and conditional field body; coordinate validation reveal and focus.", "FieldGroupDefinition, draft state, validation index", "Visibility changes never mutate values; use native details/summary or equivalent accessible behavior."],
            ["Group status selector", "Derive field count, customized count, warning count and error count for a category.", "Draft diffs, validation result", "Pure selector with no persistence or rendering side effects."],
            ["Icon", "Render the approved local icon subset beside visible labels.", "Bundled icon assets", "One family; decorative instances are aria-hidden; do not introduce a package solely for mockup decoration."],
            ["Catalogue repository", "Load and expose validated Workflow, group, field, Preset and knowledge definitions.", "Static catalogue now; replaceable adapter later", "Runtime consumers do not know whether data came from source files or a future Admin-backed service."],
            ["Catalogue validator", "Validate schema, IDs, ownership and references before publication or startup.", "Canonical catalogue schema", "Structural catalogue errors fail fast and are never repaired silently."],
        ],
        widths=[1.35, 2.25, 1.25, 1.7],
    )
    add_heading(doc, "ADR-016 — Source-authored catalogue before Admin", level=2)
    add_body(doc, "Decision: keep Workflow and Preset catalogue data in versioned TypeScript or JSON for the MVP. Expose it through CatalogueRepository and validate it at build time. Do not add an Admin route, authentication system, database, publishing workflow or generic form builder to the MVP.")
    add_bullets(doc, [
        "Reason: one researched Workflow does not justify a second product for catalogue operations.",
        "Preserved path: stable IDs, explicit schemas and repository boundaries allow a later protected Admin tool to write the same canonical representation.",
        "Admin boundary: if introduced, Admin is a separate protected route or application surface and is not shown in ordinary Workspace navigation.",
        "Form-builder boundary: begin later with ordered groups and schema-driven field definitions; add drag-and-drop or arbitrary conditional logic only if repeated Workflow authoring proves it necessary.",
    ])
    add_heading(doc, "Icon delivery decision", level=2)
    add_body(doc, "Use a small bundled SVG subset from one established MIT-licensed icon family, or an icon package already approved for another requirement. Keep accessible names in text. The Stage 6 system glyphs are visual placeholders and are not the production web-delivery strategy.")


def append_doc05(doc: DocumentType) -> None:
    add_heading(doc, "Grouped Build interaction and icon system", level=1)
    add_heading(doc, "Configuration categories", level=2)
    add_bullets(doc, [
        "Build uses a vertically ordered set of disclosure groups. Each summary includes step number, category name and compact status.",
        "Product identity and Lighting and composition open by default for the Product Workflow. Camera and optics, Materials and styling, and Output and guardrails begin collapsed.",
        "More than one group may be open. User-controlled open state remains stable during the active draft session.",
        "A structural error opens the first affected group and focuses its summary or first invalid field. Every affected field remains individually identified.",
        "An art-direction warning maps to the same group and field location, but remains confirmable and never becomes a structural error.",
        "Collapsed controls do not receive keyboard focus. Enter and Space toggle a focused group summary.",
    ])
    add_heading(doc, "Iconography", level=2)
    add_bullets(doc, [
        "Use one outlined icon family with consistent optical size and alignment.",
        "Use icons for primary navigation, Workspace and Project recognition, disclosure state and high-value actions such as Save and Generate.",
        "Keep visible text beside icons. Tooltips may supplement unfamiliar actions but never replace accessible names.",
        "Do not decorate every field label or paragraph. Text remains primary where recognition does not improve.",
    ])
    add_heading(doc, "Future Admin experience boundary", level=2)
    add_body(doc, "No Admin page is part of the MVP screen set. If catalogue editing becomes frequent enough to justify it, create a separate protected surface for trusted maintainers. It may manage Workflow metadata, ordered field groups, field definitions, Presets and preview assets through the canonical catalogue schema. It must not appear as a Project or ordinary Workspace feature.")
    add_heading(doc, "Additional UX acceptance criteria", level=2)
    add_bullets(doc, [
        "UX-AC-026: A user can identify the five configuration categories and their completion or issue status without opening each category.",
        "UX-AC-027: Validation reveals the correction location without discarding draft values or unrelated disclosure choices.",
        "UX-AC-028: Primary navigation remains understandable when icons are ignored, unavailable or announced as decorative.",
        "UX-AC-029: Workspace contains Prompt and Project work only; no catalogue-maintenance control competes with the user's task.",
    ])


def append_stage13_doc01(doc: DocumentType) -> None:
    add_heading(doc, "Selected-Preset context and MVP catalogue boundary", level=1)
    add_body(doc, "This section is authoritative where earlier mockups omit the selected Preset image, use placeholder system glyphs, vary corner treatment or imply that Workflow CRUD is needed to prove MVP CRUD capability.")
    add_table(
        doc,
        ["ID", "Decision", "Status", "Consequence"],
        [
            ["DR-017", "Build shows the selected Preset image and creative intent beside the configuration form.", "Accepted", "Users can see the direction they are remixing without leaving Build; the image remains a reference, not an output promise."],
            ["DR-018", "Use Lucide outline icons and a 7 px control / 10 px panel radius system.", "Accepted", "The earlier soft utility feel becomes reproducible without relying on operating-system glyphs."],
        ],
        widths=[0.72, 2.8, 0.8, 2.15],
    )
    add_heading(doc, "Scope consequence", level=2)
    add_bullets(doc, [
        "Projects and Saved Prompts provide authentic create, read, update and delete behavior in the user MVP.",
        "Workflow CRUD is not added merely for demonstration. A runtime Workflow is incomplete unless fields, groups, Presets, validation rules, templates and knowledge entries remain referentially valid.",
        "An under-developed Preset manager is not exposed. If portfolio needs justify it after MVP, a protected Workflow Registry Demo may clone a valid Workflow, edit metadata and order, archive it, and delete only unreferenced drafts.",
    ])


def append_stage13_doc02(doc: DocumentType) -> None:
    add_heading(doc, "Selected-Preset visual context and continuity requirements", level=1)
    add_table(
        doc,
        ["ID", "Requirement", "Priority", "Verification"],
        [
            ["FR-107", "Build shall show the active Preset sample image, name and concise creative intent beside or directly above its configuration groups.", "Must", "The selected Preset remains visually identifiable without leaving Build."],
            ["FR-108", "Changing the Preset selection shall update the visual context immediately before Generate or Remix.", "Must", "Preview identity matches activePresetId after every selection change."],
            ["FR-109", "Preset context shall include alternative text, non-guarantee wording and a deterministic missing-asset fallback.", "Must", "Missing imagery never hides the selected Preset identity or blocks configuration."],
            ["FR-110", "The interface shall use Lucide outline icons at a nominal 18 px size and 2 px stroke with visible labels for primary actions and navigation.", "Should", "Icons load locally, keep accessible names in text and do not depend on system glyphs."],
            ["FR-111", "Controls shall use a 7 px radius and cards, panels, tables and media containers shall use a 10 px radius.", "Should", "Visual review finds no unrelated corner systems on the same surface."],
            ["FR-112", "The MVP shall not expose incomplete Workflow or Preset CRUD as a demonstration feature.", "Must", "User navigation contains only functional Workspace and catalogue-consumption flows."],
        ],
        widths=[0.62, 3.05, 0.62, 2.2],
    )
    add_heading(doc, "Additional acceptance criteria", level=2)
    add_bullets(doc, [
        "AC-063: Selecting Sculptural Studio Hero displays its matching sample, name and intent before the configuration categories.",
        "AC-064: Selecting another Preset changes the preview and its alternative text without requiring Generate.",
        "AC-065: If a preview asset cannot load, Build shows a labelled fallback and retains a usable Preset selector.",
        "AC-066: Primary navigation and high-value actions use locally delivered Lucide icons with visible text labels.",
        "AC-067: Controls and large containers follow the 7 px and 10 px radius tokens across the five Stage 6 pages.",
        "AC-068: No disabled, placeholder or partially connected Workflow or Preset Admin control appears in MVP navigation.",
    ])


def append_stage13_doc03(doc: DocumentType) -> None:
    add_heading(doc, "Selected-Preset preview state", level=1)
    add_body(doc, "Preset preview is derived from the canonical catalogue and the active Build selection. It is not a second persisted copy of Preset data and does not alter the Workspace persistence model.")
    add_table(
        doc,
        ["State or value", "Source", "Lifetime", "Invariant"],
        [
            ["activePresetId", "Build draft selection", "Active draft session; saved through the existing configuration reference", "References a Preset owned by activeWorkflowId."],
            ["Selected preview", "Catalogue lookup by activePresetId", "Derived UI state", "The selectedPresetPreview contract contains name, intent, previewAssetId and accessible description; it never becomes an independently editable entity."],
            ["preview asset", "Built-in catalogue asset", "Application release or later catalogue version", "Missing assets resolve to a deterministic labelled fallback; binary data is not copied into localStorage Prompt records."],
        ],
        widths=[1.28, 1.55, 1.4, 2.25],
    )
    add_heading(doc, "Catalogue administration consequence", level=2)
    add_body(doc, "Workflow CRUD remains deferred. A future writer must validate Workflow ownership, field references, Preset defaults, rule references, template profiles and preview assets as one catalogue transaction before publication. The current localStorage design continues to store user Workspace aggregates only.")


def append_stage13_doc04(doc: DocumentType) -> None:
    add_heading(doc, "Preset context and visual-token implementation seams", level=1)
    add_table(
        doc,
        ["Component or seam", "Responsibility", "Input", "Constraint"],
        [
            ["Selected Preset preview", "Render sample image, identity, creative intent, details link, disclaimer and fallback next to Build configuration.", "activePresetId plus CatalogueRepository", "Pure view derived from catalogue state; it does not write configuration or duplicate asset data."],
            ["Icon", "Render the approved Lucide symbol beside visible text.", "Named local icon import", "Nominal 18 px, 2 px stroke; no CDN or operating-system glyph dependency."],
            ["Radius tokens", "Keep soft utility geometry consistent.", "radiusControl = 7 px; radiusPanel = 10 px", "Pills and circles are reserved for status, avatars and numbered steps."],
        ],
        widths=[1.35, 2.45, 1.45, 1.65],
    )
    add_heading(doc, "ADR-017 — Workflow CRUD remains post-MVP", level=2)
    add_body(doc, "Decision: do not add Workflow CRUD merely to demonstrate CRUD. The MVP already demonstrates complete CRUD through Projects and Saved Prompts. A disconnected editor would be false capability; a connected editor would require field, Preset, validation, template, knowledge, migration and publication integrity beyond the MVP.")
    add_bullets(doc, [
        "Accepted alternative after MVP: a protected Workflow Registry Demo that clones a valid Workflow, edits metadata and order, archives or unarchives, and deletes only unreferenced drafts.",
        "Explicitly deferred: generic field builder, Preset CRUD UI, conditional-logic designer, asset upload, role management and catalogue publishing workflow.",
        "Runtime protection: Build consumes only catalogue versions that pass schema and referential validation.",
    ])


def append_stage13_doc05(doc: DocumentType) -> None:
    add_heading(doc, "Selected Preset context and visual continuity", level=1)
    add_heading(doc, "Build behavior", level=2)
    add_bullets(doc, [
        "Place the selected Preset sample directly below the Workflow and Preset selectors and before the configuration categories on desktop.",
        "Show the sample image, Preset name, short creative intent and a details link. State visibly that the reference does not guarantee a generated output.",
        "Update the preview immediately when the Preset selection changes. The user does not need to Generate or Remix first.",
        "On narrow layouts, keep the preview in document order between the selectors and configuration categories; do not hide it behind hover.",
        "If the image is unavailable, show a stable labelled fallback while preserving the name, intent and ability to continue.",
    ])
    add_heading(doc, "Icon and shape rules", level=2)
    add_bullets(doc, [
        "Use Lucide outline icons at a nominal 18 px size and 2 px stroke for the accepted navigation and high-value actions.",
        "Keep adjacent text labels and mark decorative icon instances hidden from assistive technology.",
        "Use 7 px corners for inputs, selects, textareas and compact buttons; use 10 px corners for cards, panels, tables and media containers.",
        "Reserve circles and full pills for step markers, avatars and short statuses rather than ordinary content containers.",
    ])
    add_heading(doc, "Additional UX acceptance criteria", level=2)
    add_bullets(doc, [
        "UX-AC-030: A user can identify the selected Preset and its creative direction while viewing Build configuration.",
        "UX-AC-031: Keyboard or mouse selection changes the Preset preview without unexpected focus movement or value loss.",
        "UX-AC-032: Missing imagery leaves a labelled, understandable and operable Build state.",
        "UX-AC-033: The five Stage 6 pages use one Lucide icon treatment and one 7 px / 10 px radius system.",
        "UX-AC-034: No incomplete Workflow or Preset administration control appears in the Workspace shell.",
    ])


def append_stage14_doc01(doc: DocumentType) -> None:
    add_heading(doc, "Contextual learning, Preset recovery and selected state", level=1)
    add_body(doc, "This section records the owner-approved Stage 6 v1.4 behavior and is authoritative where earlier wording leaves Learn, reset scope, selected icons or sample-image provenance ambiguous.")
    add_table(
        doc,
        ["ID", "Decision", "Status", "Consequence"],
        [
            ["DR-019", "Learn opens compact contextual help and offers a same-tab deep link to the exact full Field Encyclopedia entry.", "Accepted", "Users get quick guidance without searching again or reading a long article inside a modal."],
            ["DR-020", "Reset to Preset restores only values controlled by the active Preset to its application snapshot.", "Accepted", "Recovery does not erase product identity, destination or generated Prompt text."],
            ["DR-021", "Persistent selection changes icon and label color while retaining non-color cues.", "Accepted", "The UI matches the preferred earlier affordance without relying on color alone."],
        ],
        widths=[0.72, 2.75, 0.8, 2.2],
    )
    add_heading(doc, "Preview-asset evidence boundary", level=2)
    add_bullets(doc, [
        "The six current Preset samples were generated with built-in AI raster image generation from direct art-direction prompts; they were not downloaded from the web.",
        "They were not produced by executing the complete DISC-01 Preset mold, field catalogue or compiler and therefore are UI direction references, not compiler-fidelity evidence.",
        "A later product-valid sample pass shall record Workflow ID, Preset ID, catalogue version, field values, compiled Prompt, generator and output asset.",
    ])


def append_stage14_doc02(doc: DocumentType) -> None:
    add_heading(doc, "Contextual learning, Preset recovery and selected-state requirements", level=1)
    add_table(
        doc,
        ["ID", "Requirement", "Priority", "Verification"],
        [
            ["FR-113", "Learn shall open a compact non-modal view for the exact configuration field, containing its name, plain-language definition and effect.", "Must", "Every Learn trigger resolves to the matching field content without a search step."],
            ["FR-114", "Open full article shall load Field Encyclopedia in the same tab at the exact field and preserve the active draft and return context.", "Must", "The route contains the field slug and Build return parameter; the draft is unchanged after return."],
            ["FR-115", "Escape, outside click and the close control shall dismiss contextual Learn and return focus to its trigger.", "Must", "Keyboard and pointer checks complete without focus loss or a modal trap."],
            ["FR-116", "A customized active Preset shall expose Reset to Preset and require confirmation before restoration.", "Must", "The dialog identifies the target and customized-field count before mutation."],
            ["FR-117", "Reset to Preset shall restore only Preset-controlled values and shall preserve product identity, Prompt name, destination, selected Preset and generated or edited Prompt text.", "Must", "Before/after state comparison confirms only the allowed configuration diff."],
            ["FR-118", "Persistent selected items shall change icon and label color and retain a non-color cue such as rail, tint or weight.", "Should", "Selected navigation and Project items remain identifiable in monochrome inspection."],
            ["FR-119", "Preset preview assets shall disclose that they are direction references and shall maintain a provenance record.", "Should", "Asset record names the generation method, exact prompt and evidence limitation."],
        ],
        widths=[0.62, 3.1, 0.62, 2.15],
    )
    add_heading(doc, "Additional acceptance criteria", level=2)
    add_bullets(doc, [
        "AC-069: Selecting Learn for Key light quality opens compact help for Key light quality and not a generic Encyclopedia landing page.",
        "AC-070: Selecting Open full article loads the exact Key light quality entry in the same tab and returning preserves the Build draft.",
        "AC-071: Pressing Escape closes compact help and restores focus to the originating Learn control.",
        "AC-072: Confirming Reset to Preset restores customized Preset-controlled values and leaves the generated Prompt text unchanged.",
        "AC-073: Cancelling Reset to Preset changes no value.",
        "AC-074: A selected Project changes its folder icon and label treatment while retaining the selected background and rail.",
        "AC-075: Every current Preset sample has a local provenance entry and non-guarantee disclosure.",
    ])


def append_stage14_doc03(doc: DocumentType) -> None:
    add_heading(doc, "Contextual-learning and Preset-reset state", level=1)
    add_table(
        doc,
        ["State or value", "Source", "Lifetime", "Invariant"],
        [
            ["activeKnowledgeFieldId", "Learn trigger or Encyclopedia route", "Navigation state", "References a knowledge entry owned by the active Workflow field; it never requires free-text search."],
            ["returnToBuild", "Learn full-article navigation", "Navigation state", "Carries route context only; the active draft remains in the existing draft store."],
            ["presetResetBaseline", "Active Preset plus catalogueVersion", "Derived draft-session state", "Contains only Preset-controlled defaults; it is recomputable and is not a standalone localStorage entity."],
            ["presetCustomizationDiff", "Current configuration compared with presetResetBaseline", "Derived UI state", "Determines reset availability and count; generated Prompt text and user-owned identity are excluded."],
        ],
        widths=[1.4, 1.7, 1.45, 2.25],
    )
    add_heading(doc, "Persistence consequence", level=2)
    add_bullets(doc, [
        "No database table or localStorage collection is added for contextual help, selected color or Reset to Preset.",
        "A saved configuration retains activePresetId, catalogueVersion and values under the existing contract; the reset baseline is derived from that versioned catalogue.",
        "If a historical catalogue version is unavailable, the application must not silently reset against a different Preset definition; it exposes an unavailable-reset state instead.",
        "Preset preview provenance remains documentation metadata and local asset evidence, not user Workspace data.",
    ])


def append_stage14_doc04(doc: DocumentType) -> None:
    add_heading(doc, "Contextual-help and Preset-recovery implementation seams", level=1)
    add_table(
        doc,
        ["Component or seam", "Responsibility", "Input", "Constraint"],
        [
            ["FieldLearnPopover", "Render compact field definition and effect; manage dismissal and focus return.", "fieldId plus KnowledgeRepository", "Non-modal; no duplicate long-form article; no extra package required."],
            ["Knowledge route", "Open the exact full article and carry Build return context.", "workflowId, fieldId, returnTo", "Preserve draft state; default to same-tab navigation while supporting native explicit new-tab behavior."],
            ["Preset reset selector", "Derive Preset-controlled baseline and current customization diff.", "activePresetId, catalogueVersion, draft configuration", "Pure TypeScript transformation; excludes product identity and Prompt text."],
            ["Reset confirmation", "Explain scope and apply the allowed diff after confirmation.", "Preset reset selector output", "No Generate, Remix, Save or Project mutation side effect."],
            ["Selected-state tokens", "Apply selected blue to label and icon with rail, tint and weight cues.", "route or selection state", "Use existing CSS tokens and Lucide icons; color is never the only cue."],
        ],
        widths=[1.35, 2.3, 1.65, 1.75],
    )
    add_heading(doc, "ADR-018 — Progressive contextual learning", level=2)
    add_body(doc, "Decision: use a compact non-modal field popover for immediate help and a same-tab exact-field Encyclopedia route for full reading. Do not place the full article in a modal and do not force a new browser tab.")
    add_heading(doc, "ADR-019 — Derived Preset reset baseline", level=2)
    add_body(doc, "Decision: derive reset state from the active Preset and catalogue version using pure TypeScript. Add no library, database entity or duplicate persisted snapshot. When the required catalogue version cannot be resolved, disable reset rather than applying a different definition.")


def append_stage14_doc05(doc: DocumentType) -> None:
    add_heading(doc, "Contextual learning, Reset to Preset and selected state", level=1)
    add_heading(doc, "Learn interaction", level=2)
    add_bullets(doc, [
        "Learn or Learn about this field opens a compact non-modal popover beside the originating field.",
        "The compact view contains the exact field name, a plain-language definition and what the choice changes. It does not duplicate the full article.",
        "Open full article uses same-tab navigation to the exact Field Encyclopedia entry and carries a return-to-Build context. The active draft is preserved.",
        "Ctrl/Cmd-click and the browser context menu may open the deep link in a new tab when explicitly requested; new-tab behavior is not forced.",
        "Escape, click outside or Close dismisses the popover and returns focus to the originating control. The popover does not trap focus or cover its field.",
    ])
    add_heading(doc, "Reset to Preset interaction", level=2)
    add_bullets(doc, [
        "Use the label Reset to Preset; do not use Reset to default because the target would be ambiguous.",
        "Enable or show the action only when a selected Preset has at least one customized Preset-controlled value.",
        "Confirmation names the selected Preset, states the customized-field count and explains what will and will not change.",
        "Confirm restores only Preset-controlled values. Preserve product identity, Prompt name, destination, selected Preset and generated or independently edited Prompt text.",
        "Cancel and Escape close the confirmation without changing values. Reset never triggers Generate, Remix or Save.",
    ])
    add_heading(doc, "Persistent selected state", level=2)
    add_bullets(doc, [
        "Selected labels and icons use the same blue. A selected Project folder may use a filled blue icon; general navigation keeps the Lucide outline form.",
        "Retain the selected left rail, background tint and stronger text weight so color is not the only cue.",
        "Treat this as persistent selection, not the brief pointer-down active state.",
    ])
    add_heading(doc, "Additional UX acceptance criteria", level=2)
    add_bullets(doc, [
        "UX-AC-035: Learn provides useful compact guidance without hiding the field or requiring a search.",
        "UX-AC-036: Full guidance opens at the exact entry, preserves the draft and offers an obvious return path.",
        "UX-AC-037: Reset confirmation prevents accidental loss and changes only Preset-controlled values.",
        "UX-AC-038: Generated Prompt text remains byte-for-byte unchanged after Reset to Preset.",
        "UX-AC-039: Selected navigation and Project state remain legible without relying on icon color alone.",
        "UX-AC-040: Preset images are labelled as illustrative references and link to a maintained provenance record.",
    ])


APPENDERS = {
    "DOC-01": append_stage14_doc01,
    "DOC-02": append_stage14_doc02,
    "DOC-03": append_stage14_doc03,
    "DOC-04": append_stage14_doc04,
    "DOC-05": append_stage14_doc05,
}


def main() -> None:
    replacements = {
        "DOC-01 v2.2": "DOC-01 v2.3",
        "DOC-02 v2.1": "DOC-02 v2.2",
        "DOC-03 v2.0": "DOC-03 v2.1",
        "DOC-04 v1.8": "DOC-04 v1.9",
        "DOC-05 v1.4": "DOC-05 v1.5",
        "DOC 01 v2.2": "DOC 01 v2.3",
        "DOC 02 v2.1": "DOC 02 v2.2",
        "DOC 03 v2.0": "DOC 03 v2.1",
        "DOC 04 v1.8": "DOC 04 v1.9",
        "DOC 05 v1.4": "DOC 05 v1.5",
        "DOC-01 Product and Project Brief v2.2": "DOC-01 Product and Project Brief v2.3",
    }
    for item in DOCS:
        doc = Document(item["src"])
        replace_in_runs(doc, {item["old"]: item["new"], **replacements})
        update_related_versions(doc)
        update_version_history(doc, item["new"], item["history"])
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        APPENDERS[item["id"]](doc)
        doc.core_properties.comments = item["history"]
        item["dst"].parent.mkdir(parents=True, exist_ok=True)
        doc.save(item["dst"])
        print(item["dst"])


if __name__ == "__main__":
    main()

