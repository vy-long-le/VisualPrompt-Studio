from pathlib import Path
from collections import Counter
from zipfile import ZipFile
import json, hashlib
from docx import Document
from lxml import etree

ROOT=Path(__file__).resolve().parents[3]
WORK=Path(__file__).resolve().parent
paths=sorted(p for p in ROOT.rglob('*.docx') if 'Previous_Versions' not in p.parts and 'Formatting_Audit_Work' not in p.parts)
def val(obj):
    if obj is None:return None
    if hasattr(obj,'pt'):return round(obj.pt,2)
    return str(obj)
def fmt(p):
    f=p.paragraph_format
    return {k:val(getattr(f,k)) for k in ['alignment','space_before','space_after','line_spacing','left_indent','right_indent','first_line_indent','keep_with_next','keep_together','page_break_before','widow_control']}
def styleinfo(s):
    f=s.font
    return dict(name=s.name,base=s.base_style.name if s.base_style else None,font=f.name,size=val(f.size),bold=f.bold,italic=f.italic,color=str(f.color.rgb),paragraph=fmt(s))
items=[]
for path in paths:
    d=Document(path)
    fonts=Counter();body=Counter();tablefonts=Counter();headings=[]
    for p in d.paragraphs:
        if p.style.name in ['Title','Subtitle'] or p.style.name.startswith('Heading'):headings.append({'text':p.text,'style':p.style.name,'direct':fmt(p),'runs':[(r.text,r.font.name,val(r.font.size),str(r.font.color.rgb),r.bold) for r in p.runs]})
        for r in p.runs:
            if r.text.strip():fonts[str((r.font.name,val(r.font.size),str(r.font.color.rgb)))]+=1
        if p.style.name=='Normal' and p.text.strip():body[str(fmt(p))]+=1
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs:
                        if r.text.strip():tablefonts[str((r.font.name,val(r.font.size),str(r.font.color.rgb)))]+=1
    sec=[]
    for s in d.sections:
        sec.append({k:round(getattr(s,k).cm,3) for k in ['page_width','page_height','top_margin','bottom_margin','left_margin','right_margin','header_distance','footer_distance']})
    with ZipFile(path) as z:
        xml=etree.fromstring(z.read('word/document.xml'));ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        tables=[]
        for t in xml.xpath('//w:tbl',namespaces=ns)[:4]:
            tables.append({'properties':etree.tostring(t.find('w:tblPr',ns),encoding='unicode'),'grid':etree.tostring(t.find('w:tblGrid',ns),encoding='unicode'),'firstRow':etree.tostring(t.find('w:tr',ns),encoding='unicode')[:5500]})
    items.append({'path':str(path.relative_to(ROOT)),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'paragraphs':len(d.paragraphs),'tablesCount':len(d.tables),'sections':sec,'styles':[styleinfo(d.styles[n]) for n in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet','List Number','Caption'] if n in d.styles],'topDirectFonts':fonts.most_common(8),'bodyDirect':body.most_common(4),'tableFonts':tablefonts.most_common(8),'front':[{'text':p.text,'style':p.style.name,'direct':fmt(p)} for p in d.paragraphs[:22]],'headings':headings,'tableSamples':tables,'headers':[p.text for s in d.sections for p in s.header.paragraphs],'footers':[p.text for s in d.sections for p in s.footer.paragraphs]})
(WORK/'current_document_formatting_audit.json').write_text(json.dumps(items,ensure_ascii=False,indent=2),encoding='utf8')
for x in items:
    print(x['path']);print('Sections:',x['sections']);print('Styles:',[(s['name'],s['font'],s['size'],s['color'],s['paragraph']['space_after'],s['paragraph']['line_spacing']) for s in x['styles']]);print('DirectFonts:',x['topDirectFonts']);print('Tables:',x['tablesCount'],x['tableFonts']);print('Front:',[(p['style'],p['text'][:85]) for p in x['front'][:10]]);print('Footer:',x['footers'])
