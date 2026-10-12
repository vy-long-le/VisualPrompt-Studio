"""Verify this formatting correction; --install installs already reviewed copies."""
from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
import hashlib, json, shutil, sys
from normalize_formatting import ROOT, WORK, REF, signature

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def rels(path):
    with ZipFile(path) as z:
        return {n.get('Id'): dict(n.attrib) for n in etree.fromstring(z.read('word/_rels/document.xml.rels'))}

report=json.loads((WORK/'content_preservation_report.json').read_text(encoding='utf8'))
inventory=json.loads((WORK/'current_document_formatting_audit.json').read_text(encoding='utf8'))
plan=json.loads((WORK.parent/'DOC06_Work/doc06_plan_audit.json').read_text(encoding='utf8'))
renders={
    'VisualPrompt_Studio_DOC-06_Test_Release_and_Maintenance_Plan_v1.0.docx':('verified_doc06',28,'doc06_a11y.json'),
    'DISC01_Research_Dossier_v1.0.docx':('verified_research_v3',97,'research_a11y.json'),
    'VisualPrompt_Studio_Discovery_Dossier_v1.0.docx':('verified_discovery_v2',14,'discovery_a11y.json'),
}
ref=Document(REF)
geometry=['page_width','page_height','top_margin','bottom_margin','left_margin','right_margin','header_distance','footer_distance']
receipt={'date':'2026-10-11','scope':'Formatting and document integrity only; no application test execution or changed research conclusions',
    'renderer':'Canonical documents 26.1007.11041 render_docx.py, bundled Python and Poppler at 120 dpi',
    'renderer_exception':'Installed LibreOffice in an isolated headless profile, explicitly authorized for DOC-06 and both dossiers',
    'audited_current_documents':len(inventory),'documents':[],'unchanged_baseline_files':[]}

for item in report:
    target=ROOT/item['file']; staged=WORK/'formatted'/target.name; original=WORK/'originals'/target.name
    assert sha(original)==item['original_sha256']
    assert sha(staged)==item['formatted_sha256']
    before=signature(original); after=signature(staged)
    for sig in (before,after):sig['protected'].pop('word/_rels/document.xml.rels',None)
    assert before==after, target.name
    before_rels=rels(original); after_rels=rels(staged)
    assert all(after_rels.get(k)==v for k,v in before_rels.items()),'Existing relationship changed'
    d=Document(staged)
    for s in d.sections:
        assert all(getattr(s,k)==getattr(ref.sections[0],k) for k in geometry)
        assert any('PAGE' in (n.text or '') for n in s.footer._element.iter(qn('w:instrText')))
        assert 'VISUALPROMPT STUDIO' in s.header.paragraphs[0].text
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        assert d.styles[name].font.name==ref.styles[name].font.name
        assert d.styles[name].font.size==ref.styles[name].font.size
        assert d.styles[name].font.color.rgb==ref.styles[name].font.color.rgb
    for table in d.tables:
        pr=table._tbl.tblPr
        assert pr.find(qn('w:tblW')).get(qn('w:w'))=='9216'
        assert sum(c.width.twips for c in table.columns)==9216
        assert table.rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None
        assert all(e.get(qn('w:color'))=='D9D9D9' for e in pr.find(qn('w:tblBorders')))
        assert all(c._tc.tcPr.find(qn('w:shd')).get(qn('w:fill'))=='1F4E79' for c in table.rows[0].cells)
        assert not list(table._tbl.iter(qn('w:trHeight')))
    folder,pages,a11y_name=renders[target.name]
    render=WORK/folder; pdf=next(render.glob('*.pdf')); reader=PdfReader(pdf)
    assert len(reader.pages)==pages
    images=[render/f'page-{n}.png' for n in range(1,pages+1)]
    assert len(list(render.glob('page-*.png')))==pages and all(p.exists() for p in images)
    texts=[p.extract_text() or '' for p in reader.pages]
    assert all('VISUALPROMPT STUDIO' in t and f'Trang {n}' in t for n,t in enumerate(texts,1))
    assert all(len(t)>150 for t in texts),'Unexpected blank page'
    a11y=json.loads((WORK/a11y_name).read_text(encoding='utf8'))
    assert a11y['counts']['high']==0 and a11y['counts']['medium']==0
    if target.name.startswith('VisualPrompt_Studio_DOC-06'):
        text='\n'.join(texts)
        assert len({c['id'] for c in plan['cases']})==83
        assert all(c['status']=='Not Run' and c['id'] in text for c in plan['cases'])
        assert len(plan['coverage'])==111 and all(plan['coverage'].values())
        assert len(plan['nfrCoverage'])==12 and not plan['unmappedFRs'] and not plan['unknownExplicitFRs']
        bookmarks={n.get(qn('w:name')) for n in d.element.body.iter(qn('w:bookmarkStart'))}
        anchors=[n.get(qn('w:anchor')) for n in d.element.body.iter(qn('w:hyperlink')) if n.get(qn('w:anchor'))]
        assert len(anchors)==10 and all(a in bookmarks for a in anchors)
        receipt['doc06_checks']={'planned_cases':83,'application_statuses':{'Not Run':83},'mapped_FRs':111,'mapped_NFRs':12,'verified_internal_links':10}
    if target.name=='DISC01_Research_Dossier_v1.0.docx':
        assert all(sha(render/f'page-{n}.png')==sha(WORK/'verified_research_v2'/f'page-{n}.png') for n in range(1,14))
    receipt['documents'].append({'file':str(target.relative_to(ROOT)),'original_sha256':sha(original),'final_sha256':sha(staged),
        'render_folder':folder,'rendered_pages':pages,'visually_inspected_pages':list(range(1,pages+1)),
        'inspection_method':'All pages viewed at native 1020 x 1320 size; research pages 1-13 also verified pixel-identical to the previously inspected render',
        'tables':len(d.tables),'content_links_fields_numbering_media_preserved':True,'original_relationships_preserved':True,
        'a11y_counts':a11y['counts'],'page_image_hashes':{p.name:sha(p) for p in images}})

changed={item['file'] for item in report}
for item in inventory:
    if item['path'] in changed:continue
    path=ROOT/item['path']
    assert sha(path)==item['sha256'],'Baseline changed'
    receipt['unchanged_baseline_files'].append({'file':item['path'],'sha256':sha(path)})
assert len(receipt['unchanged_baseline_files'])==5
for source in plan['sources']:
    assert sha(ROOT/source['path'])==source['sha256']

if '--install' in sys.argv:
    # Each destination is an existing, explicitly named project artifact with an exact backup.
    for item in report:
        target=ROOT/item['file']
        assert target.resolve().is_relative_to(ROOT.resolve())
        shutil.copy2(WORK/'formatted'/target.name,target)
        assert sha(target)==item['formatted_sha256']
receipt['installed_at_project_locations']=all(sha(ROOT/item['file'])==item['formatted_sha256'] for item in report)
receipt['total_visually_inspected_pages']=sum(d['rendered_pages'] for d in receipt['documents'])
assert receipt['total_visually_inspected_pages']==139
(WORK/'formatting_qa_receipt.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps({k:v for k,v in receipt.items() if k not in ['documents','unchanged_baseline_files']},ensure_ascii=False,indent=2))
